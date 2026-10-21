"""The asset view for every asset (iteration 2, brief §3.4): the parts of the plant × the hazards, from each cell's own
failure units, catalogue rows, value linkage and crosswalk — and the finding that the cells of one asset split the
plant differently, shown side by side instead of smoothed over.

Rows are the canonical parts of the plant, declared here per asset; every cell's own part names are kept beside them
(`parts_by_cell`), and a unit that spans several rows is listed under each with the words "spans rows". Dollars appear
only where a pinned file states them (the battery value binding; Hazard's substation reference budget pinned by the
substation flood binding); shares only where the artifact states them. Twin leads appear only where a twin crosswalk
exists (batteries, substations); the other assets say so.
"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any

from .common import ROOT, SENTENCES, fmt_pct, fmt_usd_compact, humanise, load_json, read_csv, rel, sha256
from .occurrence import hazard_repo
from .plain import TIER_SHORT, _join, label_case, singular, tier_words

try:
    import openpyxl  # type: ignore
except ImportError:  # pragma: no cover
    openpyxl = None

NO_TWIN = "No twin mapping yet — the Digital Twin's nodes have not been crosswalked to this asset's parts."
SUBSTATION_BUDGET_REL = "docs/plans/_cross_cutting/substation_conus/reference_budget.csv"
SUBSTATION_CROSSWALK_REL = "docs/plans/_cross_cutting/substation_conus/twin_crosswalk.csv"

# (row key, label, the subsystem / failure-unit words each cell uses for it)
ROWS: dict[str, list[tuple[str, str, list[str]]]] = {
    "solar": [
        ("pv_modules", "PV modules", ["PV_ARRAY", "PV array", "PV_ARRAY_MODULES_SOURCE_COHORT"]),
        ("mounting", "Mounting: trackers and racking", ["MOUNTING", "mounting"]),
        ("foundations", "Foundations", ["FOUNDATION", "foundation"]),
        ("inverters", "Inverters and combiners", ["INVERTER_SYSTEM"]),
        ("collection", "Collection cable", ["ELECTRICAL_COLLECTION"]),
        ("substation", "Transformers, switchgear and the GSU substation", ["SUBSTATION", "GSU and substation"]),
        ("scada", "SCADA and communications", ["SCADA", "SCADA and communications"]),
        ("grounding", "Grounding and lightning protection", ["GROUNDING_LIGHTNING"]),
        ("civil", "Civil works: roads, fencing, buildings", ["CIVIL_INFRA", "civil infrastructure"]),
        ("support", "Replacement support and fieldwork", ["replacement support", "SUPPORT"]),
    ],
    "wind": [
        ("rotor", "Rotor and blades", ["ROTOR_ASSEMBLY"]),
        ("tower", "Tower", ["TOWER", "REPEATED_TURBINE_SOURCE_NATIVE"]),
        ("nacelle", "Nacelle: gearbox, generator, yaw", ["NACELLE"]),
        ("foundation", "Turbine foundation", ["FOUNDATION"]),
        ("turbine_electrical", "Turbine-base and pad electrical", ["POWER_ELECTRONICS", "REPEATED_TURBINE_ELECTRICAL", "REPEATED_TURBINE_PAD_ELECTRICAL"]),
        ("collection", "Collection network", ["PLANT_COLLECTION"]),
        ("gsu", "GSU substation", ["FACILITY_GSU_SUBSTATION", "FACILITY_GSU_SUBSTATION_SOURCE_NATIVE", "SUBSTATION"]),
        ("control", "Control building, SCADA, met and O&M", ["CONTROL"]),
        ("civil", "Civil: roads, pads, buildings, fences", ["PLANT_CIVIL"]),
        ("support", "Fieldwork and transport support", ["SUPPORT"]),
    ],
    "real_estate": [
        ("house", "House (single-family, slab)", ["House"]),
        ("warehouse", "Warehouse", ["Warehouse"]),
        ("office", "Office (three storeys)", ["Office"]),
        ("retail", "Retail", ["Retail"]),
        ("contents", "Contents and inventory", ["CONTENTS"]),
    ],
}
# a unit whose component names several rows: it is listed under each of them
SPANS: dict[str, list[str]] = {
    "power conversion and collection": ["inverters", "collection"],
    "REPEATED_TURBINE": ["rotor", "tower", "nacelle"],
    "WT_TURBINE_FIRE_ASSEMBLY": ["rotor", "tower", "nacelle"],
    "WT_TURBINE_EQUIPMENT_ASSEMBLY": ["rotor", "tower", "nacelle"],
}
# wildfire × wind files every withheld unit under one subsystem word; the failure-unit id says where it belongs
BY_UNIT_TOKEN: list[tuple[str, str]] = [
    ("COLLECTION", "collection"), ("GSU", "gsu"), ("CONTROL", "control"), ("FOUNDATION", "foundation"), ("CIVIL", "civil"),
    ("SUPPORT", "support"), ("PAD", "turbine_electrical"), ("TURBINE_BASE", "turbine_electrical"), ("SCADA", "scada"),
]
RE_BUILDING = {"RE_RES_SFD_1S_SLAB": "house", "RE_IND_WAREHOUSE": "warehouse", "RE_COM_OFFICE_3S_SLAB": "office", "RE_COM_RETAIL": "retail"}
SUBSTATION_SYSTEMS = [
    ("substation_hv_yard", "345 kV yard: breakers, disconnects, bus, instrument transformers, deadends"),
    ("substation_lv_yard", "138 kV yard: breakers, disconnects, bus, instrument transformers, deadends"),
    ("substation_transformation", "Autotransformers, oil containment and firewall"),
    ("substation_controls_auxiliaries", "Relays, control house, station battery, SCADA, station service, control cable"),
    ("substation_civil_support", "Site work, access road, drainage and lightning shielding"),
]


# ─── each cell's own split ────────────────────────────────────────────────────────────────────────────────────

def _rows_for(aslug: str, fu: dict[str, Any], catalogue_part: str | None) -> list[str]:
    subsystem = str(fu.get("subsystem") or "")
    fid = str(fu.get("id") or "")
    if aslug == "real_estate":
        if subsystem == "BUILDING" and str(fu.get("component") or "") == "CONTENTS":
            return ["contents"]
        return []  # buildings come from the catalogue rows, not the failure units
    if subsystem in SPANS:
        return SPANS[subsystem]
    if fid in SPANS:
        return SPANS[fid]
    for key, label, words in ROWS.get(aslug, []):
        if subsystem in words:
            return [key]
    if subsystem in ("WITHHELD_OR_SUPPORT", "PLANT_ELECTRICAL", "SUPPORT", "BESS_MV_AND_SHARED_ELECTRICAL", "BESS_CIVIL", "BESS_FIELDWORK") or not subsystem:
        for token, key in BY_UNIT_TOKEN:
            if token in fid:
                return [key]
    return []


def _flood_solar_shares(folder: Path | None) -> dict[str, float]:
    """The Flood × Solar workbook's Value_Link sheet is the only place the cell states its parts' value shares."""
    if openpyxl is None or folder is None:
        return {}
    books = sorted(folder.glob("*.xlsx"))
    if not books:
        return {}
    wb = openpyxl.load_workbook(books[-1], read_only=True, data_only=True)
    if "Value_Link" not in wb.sheetnames:
        return {}
    out: dict[str, float] = {}
    header: list[str] | None = None
    for r in wb["Value_Link"].iter_rows(values_only=True):
        cells = ["" if v is None else str(v).strip() for v in r]
        if header is None:
            if "value_id" in cells and "physical_share" in cells:
                header = cells
            continue
        row = dict(zip(header, cells))
        try:
            out[row["value_id"]] = float(row["physical_share"])
        except (KeyError, ValueError):
            continue
    return out


def unit_words(fu: dict[str, Any]) -> str:
    """'FW_ELEVATED_TURBINE_EQUIPMENT' → 'elevated turbine equipment'; a CONTENTS component → 'contents and inventory'."""
    if str(fu.get("component") or "") == "CONTENTS":
        return "contents and inventory"
    fid = str(fu.get("id") or "")
    fid = re.sub(r"^[A-Z]{2,5}\d?_(?=[A-Z])", "", fid) if not fid.startswith(("SUPPORT_", "BESS_")) else fid.replace("BESS_", "")
    return humanise(fid).lower()


def cell_parts(e: Any, cell: dict[str, Any], aslug: str, bess_values: dict[str, float] | None, fb_binding: dict[str, Any] | None) -> dict[str, Any]:
    """Every failure unit and catalogue row of one cell, each mapped to the asset's canonical rows, with its share where stated."""
    a = e.artifact or {}
    vl = a.get("value_linkage") or {}
    parts: list[dict[str, Any]] = []
    basis = None
    shares: dict[str, float] = {}
    share_basis = None
    if e.id == "hail_solar":
        prof = (vl.get("value_profiles") or [{}])[0]
        shares = {prof.get("failure_unit_id", ""): float(prof.get("failure_unit_share_physical_base", 0) or 0)}
        share_basis = "share of the physical replaceable base (direct module hardware profile)"
    elif e.id == "tropical_cyclone_wind_solar":
        prof = vl.get("screening_reference_profile") or {}
        denom = float(prof.get("physical_replacement_value_usd_per_kwdc") or 0) or None
        if denom:
            shares = {k: float(v) / denom for k, v in (prof.get("unit_values") or {}).items()}
            share_basis = "share of the physical replacement value, 2024 USD per kWdc (named reference profile)"
    elif e.id == "flood_solar":
        shares = _flood_solar_shares(e.folder)
        share_basis = "share of the physical replaceable base (workbook Value_Link sheet)"
    elif e.id == "tropical_cyclone_wind_wind":
        shares = {"WT_JAIMES_TURBINE_TOWER_EXPOSURE_UNIT": 0.16}
        share_basis = "share of project TIV (owner-approved tower proxy basis)"
    elif e.id in ("flood_bess", "tropical_cyclone_wind_bess", "wildfire_bess") and bess_values:
        total = sum(bess_values.values())
        share_basis = "share of the reference plant ($109.8M, reference budget r1)"
    fus = a.get("failure_units") or []
    if aslug == "real_estate":
        fus = [f for f in fus if str(f.get("component") or "") == "CONTENTS"]
    if e.kind == "binding" and aslug == "substation":
        # a binding names no failure units of its own: its catalogue rows are its parts, mapped through the ledger where one exists
        ledger = (a.get("value") or {}).get("ledger") or []
        modeled_rows = {l.get("row_id") for l in ledger if str(l.get("role", "")).startswith("modeled")}
        for r in e.data.get("rows") or []:
            covered = [k for k, _ in SUBSTATION_SYSTEMS]
            if modeled_rows:
                budget_rows = read_csv(ROOT / "dashboard" / "data" / "substation_reference_budget.csv") if (ROOT / "dashboard" / "data" / "substation_reference_budget.csv").exists() else []
                sys_of = {b["row_id"]: b["system"] for b in budget_rows}
                covered = sorted({sys_of[i] for i in modeled_rows if i in sys_of}, key=lambda k: [x for x, _ in SUBSTATION_SYSTEMS].index(k))
            parts.append({"failure_unit_id": a.get("failure_unit_id"), "subsystem": "FACILITY_GSU_SUBSTATION_SOURCE_NATIVE", "component": r["part"], "treatment": "reused",
                          "part": r["part"], "state": "curve", "rows": covered, "spans": False, "share": None, "share_basis": None, "dollars": None,
                          "shape": r.get("shape"), "cost": r.get("cost"), "curves": r.get("curves") or [], "drawn": True, "withheld_reasons": [], "whole_station": not modeled_rows,
                          "modeled_row_ids": sorted(modeled_rows)})
            if modeled_rows:
                rest = [k for k, _ in SUBSTATION_SYSTEMS if k not in covered]
                parts.append({"failure_unit_id": None, "subsystem": "FACILITY_GSU_SUBSTATION", "component": "rest of the fenced station", "treatment": "withheld",
                              "part": "rest of the fenced station", "state": "withheld", "rows": rest, "spans": False, "share": None, "share_basis": None, "dollars": None,
                              "shape": None, "cost": None, "curves": [], "drawn": False, "withheld_reasons": ["withheld by the binding"]})
    for fu in fus:
        sh = shares.get(fu.get("id")) if shares else None
        if sh is None and fu.get("default_value_share_physical_base") is not None:
            sh = float(fu["default_value_share_physical_base"])
            share_basis = share_basis or "share of the physical replaceable base (artifact default)"
        if sh is None and fu.get("direct_hardware_value_share_physical_base") is not None:
            sh = float(fu["direct_hardware_value_share_physical_base"])
        treatment = str(fu.get("treatment") or "")
        state = "withheld" if treatment.startswith("withheld") else ("exposure_modifier" if "exposure_modifier" in treatment else "curve")
        rows = _rows_for(aslug, fu, None)
        if aslug == "bess":
            rows = [fu.get("subsystem")] if fu.get("subsystem") in ("battery_enclosures", "pcs") else [{"BESS_MV_AND_SHARED_ELECTRICAL": "electrical_balance", "BESS_CIVIL": "civil_works", "BESS_FIELDWORK": "fieldwork"}.get(str(fu.get("id")), "")]
            rows = [r for r in rows if r]
            if bess_values and state == "curve":
                vb = fb_binding or {}
                if e.id == "flood_bess":
                    u = (vb.get("units") or {}).get(fu.get("subsystem")) or {}
                    dollars = float(u.get("count", 0)) * float(u.get("same_unit_value_usd", 0)) if u else None
                else:
                    other = load_json(ROOT / vl["value_binding_path"]) if vl.get("value_binding_path") else {}
                    u = (other.get("units") or {}).get(fu.get("id")) or {}
                    dollars = float(u.get("count", 0)) * float(u.get("value_usd_each", 0)) if u else None
                if dollars is not None:
                    fu = {**fu, "_dollars": dollars}
                    sh = dollars / total if total else None
        catalogue_row = next((r for r in e.data.get("rows") or [] if any(c in (r.get("curves") or []) for c in [c["curve_id"] for c in cell.get("_curves", []) if c.get("failure_unit_id") == fu.get("id")])), None)
        if aslug == "substation" and e.kind == "cell":
            rows = [k for k, _ in SUBSTATION_SYSTEMS]  # the Hazus assembly is the whole fenced station
        parts.append({"failure_unit_id": fu.get("id"), "subsystem": fu.get("subsystem"), "component": fu.get("component"), "treatment": treatment,
                      "part": (catalogue_row or {}).get("part") or unit_words(fu), "state": state, "rows": rows,
                      "spans": len(rows) > 1, "share": sh, "share_basis": share_basis if sh is not None else None, "dollars": fu.get("_dollars"),
                      "shape": (catalogue_row or {}).get("shape"), "cost": (catalogue_row or {}).get("cost"),
                      "curves": [c["curve_id"] for c in cell.get("_curves", []) if c.get("failure_unit_id") == fu.get("id")],
                      "drawn": any(c.get("drawn") for c in cell.get("_curves", []) if c.get("failure_unit_id") == fu.get("id")),
                      "withheld_reasons": fu.get("withheld_reason_codes") or []})
    if aslug == "real_estate":
        # the buildings are the catalogue rows; hail splits each into roof and walls with a share of the building's value
        for r in e.data.get("rows") or []:
            name = r["part"]
            key = next((k for k, lab, words in ROWS["real_estate"] if name.startswith(words[0])), None)
            fu_id = None
            share = None
            for fu in a.get("failure_units") or []:
                shares_rcn = fu.get("share_of_structure_rcn") or {}
                bid = next((b for b, k2 in RE_BUILDING.items() if k2 == key), None)
                comp = str(fu.get("component") or "")
                if bid in shares_rcn and ((comp == "ROOF_COVER" and "roof" in name.lower()) or (comp == "WALLS_AND_WINDOWS" and ("wall" in name.lower() or "window" in name.lower()))):
                    fu_id, share = fu.get("id"), float(shares_rcn[bid])
            sub = name[len(ROWS["real_estate"][[k for k, _, _ in ROWS["real_estate"]].index(key)][2][0]):].strip() if key else ""
            sub = re.sub(r"\s*\([^)]*\)", "", sub).strip()
            parts.append({"failure_unit_id": fu_id, "subsystem": "BUILDING", "component": sub or "whole building", "treatment": "primary_nonzero", "part": name,
                          "state": "curve", "rows": [key] if key else [], "spans": False, "share": share,
                          "share_basis": "share of the building's structure replacement cost new" if share is not None else None, "dollars": None,
                          "shape": r.get("shape"), "cost": r.get("cost"), "curves": r.get("curves") or [], "drawn": True, "withheld_reasons": [], "whole_building": not sub})
    return {"entry_id": e.id, "hazard": cell["hazard"], "status": e.status, "title": cell["title"], "parts": parts, "share_basis": share_basis,
            "n_named": len(parts), "n_curve": sum(1 for p in parts if p["state"] == "curve"), "n_withheld": sum(1 for p in parts if p["state"] == "withheld"),
            "not_covered": [re.sub(r"\s*\(\$[\d.,]+\s*[MBk]?\)\s*$", "", s) for s in e.data.get("not_covered") or []]}


# ─── substation: Hazard's reference budget and the bindings' ledgers ──────────────────────────────────────────

def substation_budget(pin_sha: str | None, out_dir: Path) -> dict[str, Any]:
    """Hazard's substation reference budget, read from the Hazard checkout (never imported) and checked against the SHA
    the substation flood binding pins; a copy is kept beside the data so the build runs without Hazard, and says so."""
    repo = hazard_repo()
    copy = out_dir / "substation_reference_budget.csv"
    if repo is not None and (repo / SUBSTATION_BUDGET_REL).exists():
        path, source = repo / SUBSTATION_BUDGET_REL, f"Hazard_modeling/{SUBSTATION_BUDGET_REL}"
    elif copy.exists():
        path, source = copy, "dashboard/data/substation_reference_budget.csv (Hazard_modeling not found on this machine)"
    else:
        return {"available": False, "ok": None, "rows": [], "source": None, "sha256": None, "expected_sha256": pin_sha}
    rows = read_csv(path)
    actual = sha256(path)
    return {"available": True, "ok": (actual == pin_sha) if pin_sha else None, "rows": rows, "source": source, "sha256": actual, "expected_sha256": pin_sha,
            "text": path.read_text(encoding="utf-8")}


def substation_crosswalk(out_dir: Path) -> dict[str, Any]:
    repo = hazard_repo()
    copy = out_dir / "substation_twin_crosswalk.csv"
    if repo is not None and (repo / SUBSTATION_CROSSWALK_REL).exists():
        path, source = repo / SUBSTATION_CROSSWALK_REL, f"Hazard_modeling/{SUBSTATION_CROSSWALK_REL}"
    elif copy.exists():
        path, source = copy, "dashboard/data/substation_twin_crosswalk.csv (Hazard_modeling not found on this machine)"
    else:
        return {"available": False, "rows": [], "source": None, "sha256": None}
    return {"available": True, "rows": read_csv(path), "source": source, "sha256": sha256(path), "text": path.read_text(encoding="utf-8"),
            "pin": "read, not pinned by any artifact in this repository"}


# ─── the asset view ───────────────────────────────────────────────────────────────────────────────────────────

def _state_words(state: str, tier_span: str | None, extra: str | None = None) -> str:
    words = {
        "curve": "curve · " + tier_words(tier_span) if tier_span and tier_span != "–" else "curve",
        "partial": "curve for part of it · " + tier_words(tier_span) if tier_span and tier_span != "–" else "curve for part of it",
        "whole": "one curve for the whole " + (extra or "asset"),
        "reused": "reused curve · " + tier_words(tier_span) if tier_span and tier_span != "–" else "reused curve",
        "withheld": "named, no response yet (withheld)",
        "no_curve": "no curve",
        "not_named": "not named by this cell",
        "gap": "no cell yet",
        "exposure_modifier": "no own curve; takes the other parts' damage ratio",
    }
    return words.get(state, state)


def _span(codes: list[str | None]) -> str:
    tiers = sorted({c for c in codes if c})
    if not tiers:
        return "–"
    return tiers[0] if len(tiers) == 1 else f"{tiers[0]}–{tiers[-1][1]}"


def asset_view(aslug: str, a_def: dict[str, Any], entries: list[Any], cells: dict[str, dict[str, Any]], sl: dict[str, Any], registers: dict[tuple[str, str], Any],
               bess_values: dict[str, float] | None, fb_binding: dict[str, Any] | None, twin: dict[str, Any], sub_budget: dict[str, Any] | None,
               sub_crosswalk: dict[str, Any] | None, resolve_state: Any) -> dict[str, Any]:
    hz = [h for h in sl["hazards"]]
    by_hazard: dict[str, Any] = {}
    for e in entries:
        by_hazard[next(h["slug"] for h in hz if h["catalogue"] == e.data["hazard"])] = e
    parts_by_cell = {e.id: cell_parts(e, cells[e.id], aslug, bess_values, fb_binding) for e in entries if e.kind != "gap"}
    # ── rows
    rows: list[dict[str, Any]] = []
    if aslug == "bess":
        total = sum(bess_values.values()) if bess_values else None
        for b in sl["value_buckets"]["bess"]["buckets"]:
            v = (bess_values or {}).get(b["key"])
            rows.append({"key": b["key"], "label": b["label"], "value_usd": v, "share": (v / total) if v is not None and total else None, "twin_targets": b["twin_targets"], "subsystems": b["failure_unit_subsystems"], "cells": {}})
    elif aslug == "substation" and sub_budget and sub_budget.get("available"):
        by_sys: dict[str, float] = {}
        for r in sub_budget["rows"]:
            try:
                by_sys[r["system"]] = by_sys.get(r["system"], 0.0) + float(r["amount_usd"])
            except ValueError:
                pass
        total = sum(by_sys.values())
        for key, label in SUBSTATION_SYSTEMS:
            v = by_sys.get(key)
            rows.append({"key": key, "label": label, "value_usd": v, "share": (v / total) if v is not None and total else None, "twin_targets": [f"phy:substation:{key}"], "subsystems": [key], "cells": {}})
    else:
        for key, label, words in ROWS.get(aslug, []):
            rows.append({"key": key, "label": label, "value_usd": None, "share": None, "twin_targets": [], "subsystems": words, "cells": {}})
    # ── cells
    whole_word = {"substation": "fenced station", "real_estate": "building"}.get(aslug, "asset")
    for row in rows:
        for h in hz:
            e = by_hazard.get(h["slug"]) if h["catalogue"] else None
            reg = registers.get((aslug, h["slug"]))
            bucket_reg = (reg or {}).get("register", {}).get("buckets", {}).get(row["key"]) if reg else None
            occ = resolve_state(bucket_reg, reg["register"].get("surfaces", []), reg["register"].get("last_searched")) | {"bucket": row["key"]} if bucket_reg else None
            cellv: dict[str, Any] = {"entry_id": e.id if e else None, "effect": h["effect"], "outside_the_five": h["outside_the_five"],
                                     "twin_lead": (bucket_reg or {}).get("twin_lead", {}).get("level") if bucket_reg else None,
                                     "twin_nodes": (bucket_reg or {}).get("twin_lead", {}).get("nodes", []) if bucket_reg else [],
                                     "occurrence": occ, "state": None, "words": "—", "mark": "—", "parts": [], "tier_span": None, "share": None, "dollars": None}
            if e is None:
                row["cells"][h["slug"]] = cellv
                continue
            if e.kind == "gap":
                cellv.update({"state": "gap", "words": _state_words("gap", None), "mark": "┈"})
                row["cells"][h["slug"]] = cellv
                continue
            pc = parts_by_cell[e.id]
            mine = [p for p in pc["parts"] if row["key"] in p["rows"]]
            if e.kind == "binding":
                src_cell = cells[e.id]
                tier = (_span([r.get("shape") for r in e.data.get("rows") or []]), _span([r.get("cost") for r in e.data.get("rows") or []]))
                span = _span([tier[0], tier[1]]) if tier[0] != "–" else "–"
                if aslug == "substation" and e.id.startswith("substation_wildfire"):
                    ledger = (e.artifact or {}).get("value", {}).get("ledger") or []
                    sys_of = {r["row_id"]: r["system"] for r in (sub_budget or {}).get("rows", [])}
                    rows_here = [l for l in ledger if sys_of.get(l.get("row_id")) == row["key"]]
                    modeled = [l for l in rows_here if str(l.get("role", "")).startswith("modeled")]
                    if rows_here and len(modeled) == len(rows_here):
                        st, words = "reused", _state_words("reused", span)
                    elif modeled:
                        st, words = "partial", f"curve for part of it: {_join([l['item'].lower() for l in modeled])} · {tier_words(span)}"
                    else:
                        st, words = "withheld", "in the station's value, withheld by this binding"
                    cellv.update({"state": st, "words": words, "mark": "▨" if st != "withheld" else "▢", "tier_span": span,
                                  "dollars": sum(float(l["amount_usd"]) for l in modeled) if modeled else 0.0, "parts": [l["item"] for l in rows_here]})
                else:
                    cellv.update({"state": "whole", "words": _state_words("whole", span, whole_word) + f" · {tier_words(span)}", "mark": "▨", "tier_span": span,
                                  "parts": [r["part"] for r in e.data.get("rows") or []], "dollars": row.get("value_usd")})
                row["cells"][h["slug"]] = cellv
                continue
            if aslug == "substation" or (aslug == "real_estate" and row["key"] != "contents"):
                whole = [p for p in pc["parts"] if p.get("whole_building")] if aslug == "real_estate" else [p for p in pc["parts"] if p["state"] == "curve"]
                if aslug == "real_estate":
                    mine = [p for p in pc["parts"] if row["key"] in p["rows"]]
                    whole = [p for p in mine if p.get("whole_building")]
                if whole and not [p for p in mine if not p.get("whole_building")] and aslug == "real_estate":
                    span = _span([whole[0].get("shape"), whole[0].get("cost")])
                    cellv.update({"state": "whole", "words": _state_words("whole", span, whole_word) + f" · {tier_words(span)}", "mark": "■", "tier_span": span, "parts": [p["part"] for p in whole]})
                    row["cells"][h["slug"]] = cellv
                    continue
                if aslug == "substation" and whole:
                    span = _span([whole[0].get("shape"), whole[0].get("cost")])
                    cellv.update({"state": "whole", "words": _state_words("whole", span, whole_word) + f" · {tier_words(span)}", "mark": "■", "tier_span": span, "parts": [p["part"] for p in whole], "dollars": row.get("value_usd")})
                    row["cells"][h["slug"]] = cellv
                    continue
            if mine:
                curves = [p for p in mine if p["state"] == "curve"]
                withheld = [p for p in mine if p["state"] == "withheld"]
                mods = [p for p in mine if p["state"] == "exposure_modifier"]
                span = _span([x for p in curves for x in (p.get("shape"), p.get("cost"))])
                if curves:
                    dollars = sum(p["dollars"] for p in curves if p.get("dollars") is not None) if any(p.get("dollars") is not None for p in curves) else None
                    partial = (aslug == "bess" and row.get("value_usd") is not None and dollars is not None and dollars < row["value_usd"] - 1.0) or (bool(withheld) and not curves) or (aslug == "real_estate" and all(not p.get("whole_building") for p in curves))
                    st = "partial" if partial else "curve"
                    words = _state_words(st, span)
                    if aslug == "real_estate" and st == "partial":
                        words = f"curve for {_join([str(p['component']).lower() for p in curves])} only · {tier_words(span)}"
                    cellv.update({"state": st, "words": words, "mark": "■" + ("◔" if partial else ""), "tier_span": span, "parts": [p["part"] for p in mine],
                                  "share": sum(p["share"] for p in curves if p.get("share") is not None) if any(p.get("share") is not None for p in curves) else None, "dollars": dollars,
                                  "spans": [p["part"] for p in mine if p["spans"]], "drawn": any(p["drawn"] for p in curves)})
                elif withheld:
                    cellv.update({"state": "withheld", "words": _state_words("withheld", None), "mark": "▢", "parts": [p["part"] for p in mine], "spans": [p["part"] for p in mine if p["spans"]]})
                elif mods:
                    cellv.update({"state": "exposure_modifier", "words": _state_words("exposure_modifier", None), "mark": "–", "parts": [p["part"] for p in mine]})
            else:
                # not a named unit: is it in the cell's "no curve yet" list, or simply not mentioned?
                nc = pc["not_covered"]
                label_words = set(re.split(r"[^a-z]+", row["label"].lower())) - {"and", "the", "of", ""}
                hit = next((s for s in nc if label_words & (set(re.split(r"[^a-z]+", s.lower())) - {"and", "the", "of", ""})), None)
                if aslug == "bess" and e.id == "flood_bess" and row["key"] in ("electrical_balance", "civil_works", "fieldwork"):
                    hit = hit or row["label"]
                if hit:
                    cellv.update({"state": "no_curve", "words": _state_words("no_curve", None), "mark": "✕", "no_curve_text": hit,
                                  "sentence": SENTENCES["no_curve"]["page"].format(part=row["label"], value=fmt_usd_compact(row["value_usd"]) if row.get("value_usd") is not None else "an unstated share")})
                else:
                    cellv.update({"state": "not_named", "words": _state_words("not_named", None), "mark": "—"})
            row["cells"][h["slug"]] = cellv
    # ── the split-differently finding
    split = split_finding(aslug, a_def, rows, parts_by_cell, hz, sub_budget)
    # ── twin
    if aslug == "bess":
        twin_block = {"available": True, "sentence": f"Twin leads come from the Digital Twin through Hazard's crosswalk ({twin.get('rows')} rows, pinned by the battery value binding); a lead is design exposure, never evidence.",
                      "sha256": twin.get("sha256"), "rows": twin.get("rows"), "nodes": {}}
    elif aslug == "substation" and sub_crosswalk and sub_crosswalk.get("available"):
        nodes: dict[str, list[dict[str, str]]] = {}
        for r in sub_crosswalk["rows"]:
            tgt = r.get("candidate_target", "")
            key = tgt.split(":")[-1] if tgt.startswith("phy:substation:") else None
            if key:
                nodes.setdefault(key, []).append({"node": r.get("source_node_id", ""), "name": r.get("source_name", ""), "disposition": r.get("disposition", ""), "why": r.get("reason", "")})
        twin_block = {"available": True, "sentence": f"The Digital Twin's substation nodes are crosswalked to these rows ({len(sub_crosswalk['rows'])} rows, read from the Hazard repository, not pinned by any artifact here); no hazard leads are in a register for substations yet, so no lead letter is shown.",
                      "sha256": sub_crosswalk.get("sha256"), "rows": len(sub_crosswalk["rows"]), "nodes": nodes, "leads": False}
    else:
        twin_block = {"available": False, "sentence": NO_TWIN, "sha256": None, "rows": None, "nodes": {}}
    return {"rows": rows, "parts_by_cell": parts_by_cell, "split": split, "twin": twin_block,
            "value_source": ({"label": "reference budget r1 (battery value binding)", "path": "docs/cells/flood_bess/current/value_binding.json"} if aslug == "bess"
                             else {"label": "Hazard's substation reference budget, pinned by the substation flood binding", "path": SUBSTATION_BUDGET_REL, "sha256": (sub_budget or {}).get("sha256"), "pin_ok": (sub_budget or {}).get("ok")} if aslug == "substation" and sub_budget and sub_budget.get("available")
                             else None)}


def split_finding(aslug: str, a_def: dict[str, Any], rows: list[dict[str, Any]], parts_by_cell: dict[str, Any], hz: list[dict[str, Any]], sub_budget: dict[str, Any] | None) -> dict[str, Any]:
    cols = [pc for pc in parts_by_cell.values()]
    if not cols:
        return {"differs": False, "sentence": "No cell names parts for this asset yet.", "columns": [], "table": [], "details": []}
    details: list[str] = []
    counts = {pc["hazard"]["label"].lower(): pc["n_named"] for pc in cols}
    if len(set(counts.values())) > 1:
        details.append("the cells name different numbers of parts: " + _join([f"{h} {n}" for h, n in counts.items()]))
    # rows named by some cells and not others
    for row in rows:
        named = [pc["hazard"]["label"].lower() for pc in cols if any(row["key"] in p["rows"] for p in pc["parts"])]
        listed = [pc["hazard"]["label"].lower() for pc in cols if pc["hazard"]["label"].lower() not in named and row["cells"].get(next(h["slug"] for h in hz if h["label"] == pc["hazard"]["label"]), {}).get("state") == "no_curve"]
        absent = [pc["hazard"]["label"].lower() for pc in cols if pc["hazard"]["label"].lower() not in named and pc["hazard"]["label"].lower() not in listed]
        if named and (listed or absent) and aslug not in ("real_estate", "substation"):
            d = f"{label_case(row['label'])} is a unit under {_join(named)}"
            if listed:
                d += f"; {_join(listed)} list{'s' if len(listed) == 1 else ''} it as no curve yet"
            if absent:
                d += f"; {_join(absent)} {'does' if len(absent) == 1 else 'do'} not mention it"
            details.append(d)
    # units that span rows
    for pc in cols:
        for p in pc["parts"]:
            if p["spans"]:
                details.append(f"{pc['hazard']['label'].lower()} treats {p['part'].lower()} as one unit that spans {_join([next(r['label'].lower() for r in rows if r['key'] == k) for k in p['rows']])}")
    # shares that disagree for the same row
    for row in rows:
        stated = []
        for pc in cols:
            sh = [p["share"] for p in pc["parts"] if row["key"] in p["rows"] and p.get("share") is not None and p["state"] == "curve"]
            dollars = [p["dollars"] for p in pc["parts"] if row["key"] in p["rows"] and p.get("dollars") is not None and p["state"] == "curve"]
            if dollars:
                stated.append((pc["hazard"]["label"].lower(), sum(dollars), "usd"))
            elif sh:
                stated.append((pc["hazard"]["label"].lower(), sum(sh), "share"))
        if len(stated) >= 2:
            kinds = {k for _, _, k in stated}
            if "usd" in kinds and len({round(v / 1e5) for _, v, k in stated if k == "usd"}) > 1:
                details.append(f"{label_case(row['label'])}: {_join([f'{h} covers {fmt_usd_compact(v)}' for h, v, k in stated if k == 'usd'])}")
            elif "share" in kinds and (max(v for _, v, _ in stated) - min(v for _, v, _ in stated)) > 0.02:
                details.append(f"{label_case(row['label'])}: {_join([f'{h} gives it {fmt_pct(v)}' for h, v, k in stated if k == 'share'])}" + " of the plant's value")
    bases = sorted({pc["share_basis"] for pc in cols if pc.get("share_basis")})
    if len(bases) > 1:
        details.append("the shares are not on one basis: " + "; ".join(bases))
    # real estate: hail splits buildings; retail only in some
    if aslug == "real_estate":
        for pc in cols:
            subs = [p for p in pc["parts"] if p.get("component") not in ("whole building", "CONTENTS") and p["state"] == "curve"]
            if subs:
                kinds = []
                for comp in sorted({str(p["component"]).lower() for p in subs}):
                    word = "the roof" if comp.startswith("roof") else ("the walls and windows" if "wall" in comp or "window" in comp else comp)
                    if word not in kinds:
                        kinds.append(word)
                details.append(f"{pc['hazard']['label'].lower()} splits each building into {_join(kinds)}, each a share of the building's value; the other hazards price the building as one")
        named_bldg = {pc["hazard"]["label"].lower(): sorted({r_["label"].split(" (")[0].lower() for r_ in rows if r_["key"] != "contents" and any(r_["key"] in p["rows"] for p in pc["parts"])}) for pc in cols}
        all_b = sorted({b for v in named_bldg.values() for b in v})
        for b in all_b:
            without = [h for h, v in named_bldg.items() if b not in v]
            if without:
                details.append(f"{b} has a curve under {_join([h for h, v in named_bldg.items() if b in v])} only; {_join(without)} {'does' if len(without) == 1 else 'do'} not name it")
    if aslug == "substation" and sub_budget and sub_budget.get("available"):
        sys_of = {r["row_id"]: r["system"] for r in sub_budget["rows"]}
        amount = {r["row_id"]: float(r["amount_usd"]) for r in sub_budget["rows"] if r.get("amount_usd")}
        for pc in cols:
            ids = next((p.get("modeled_row_ids") for p in pc["parts"] if p.get("modeled_row_ids")), None)
            if ids:
                bound = sum(amount.get(i_, 0.0) for i_ in ids)
                fenced = sum(v for k_, v in amount.items() if sys_of.get(k_) and k_ != "road")
                details.append(f"{pc['hazard']['label'].lower()} binds {len(ids)} equipment rows ({fmt_usd_compact(bound)}, {fmt_pct(bound / fenced) if fenced else '—'} of the fenced station) and withholds the rest; the other hazards use one curve for the whole station")
    no_budget = {"solar": "No reference plant budget exists for solar farms in this repository (Hazard has one for batteries and substations); each cell states its own shares, on its own basis.",
                 "wind": "No reference plant budget exists for wind farms in this repository; the hurricane cell carries the only plant value (the canonical 100 MW farm, $140M, of which the tower share is 0.16).",
                 "real_estate": "Buildings carry no dollar value here: every curve is a share of the building's structure replacement cost, which the consumer supplies.",
                 "substation": "Hazard's substation reference budget ($27.35M fenced station) sets the rows; three hazards use one curve for the whole station, wildfire binds the equipment rows only.",
                 "bess": "Hazard's battery reference budget ($109.8M) sets the rows."}
    table = []
    for row in rows:
        table.append({"key": row["key"], "label": row["label"], "by_cell": {pc["entry_id"]: [p["part"] + (" (withheld)" if p["state"] == "withheld" else "") for p in pc["parts"] if row["key"] in p["rows"]] for pc in cols}})
    differs = bool(details)
    lead = f"These cells split the {singular(a_def.get('noun', a_def['label']))} differently" if differs else f"The cells split the {singular(a_def.get('noun', a_def['label']))} the same way"
    sentence = lead + (": " + "; ".join(details[:6]) + ("; and more below" if len(details) > 6 else "") + ". " if details else ". ") + no_budget.get(aslug, "")
    return {"differs": differs, "sentence": sentence, "details": details, "columns": [{"entry_id": pc["entry_id"], "hazard": pc["hazard"], "n_named": pc["n_named"], "n_curve": pc["n_curve"], "n_withheld": pc["n_withheld"], "share_basis": pc.get("share_basis"), "title": pc["title"]} for pc in cols], "table": table}


def asset_plain_sentence(a_def: dict[str, Any], rows: list[dict[str, Any]], hz: list[dict[str, Any]], entries_status: dict[str, str]) -> str:
    """'Hail is the one hazard with a lab-tested curve (modules); flood … Four parts of the plant have no curve under any hazard.'"""
    noun = singular(a_def.get("noun", a_def["label"]))
    five = [h for h in hz if not h["outside_the_five"]]
    groups: dict[str, list[tuple[str, str]]] = {}  # evidence words → [(hazard, what)]
    gaps: list[str] = []
    quals: list[str] = []
    for h in five:
        states = [row["cells"].get(h["slug"], {}) for row in rows]
        curves = [c for c in states if c.get("state") in ("curve", "partial", "whole", "reused")]
        status = entries_status.get(h["slug"])
        if status == "gap":
            gaps.append(h["label"].lower())
            continue
        if not curves:
            groups.setdefault("named parts but no drawn curve", []).append((h["label"].lower(), ""))
            continue
        tiers = sorted({t for c in curves for t in (c.get("tier_span") or "").replace("–", "–T").split("–") if t.startswith("T")})
        words = tier_words(f"{tiers[0]}–{tiers[-1][1]}" if len(tiers) > 1 else (tiers[0] if tiers else None))
        parts = [label_case(row["label"].split(":")[0]) for row in rows if row["cells"].get(h["slug"], {}).get("state") in ("curve", "partial")]
        if all(c.get("state") in ("whole", "reused") for c in curves):
            what = f"one curve for the whole {noun}"
        elif len(parts) == 1:
            what = f"{parts[0]} only"
        else:
            what = f"{len(parts)} parts"
        quals.append("proposal" if status == "proposal" else ("reused" if status == "reused" else "live"))
        qual = "" if len(set(quals)) <= 1 else (" · proposal" if status == "proposal" else (" · reused" if status == "reused" else ""))
        groups.setdefault(words, []).append((h["label"].lower(), what + qual))
    order = ["lab-tested", "field-calibrated", "transferred", "lab-tested and transferred", "transferred and our judgment", "our judgment"]
    keys = sorted(groups, key=lambda k_: (order.index(k_) if k_ in order else 50, k_))
    bits = []
    for k_ in keys:
        hs = groups[k_]
        names = _join([h for h, _ in hs])
        verb = "has" if len(hs) == 1 else "have"
        evidence = {"lab-tested": "lab-tested curves", "our judgment": "curves that are our judgment", "transferred": "transferred curves"}.get(k_, f"curves built from {k_} evidence")
        if k_ == "named parts but no drawn curve":
            evidence = "named parts but no drawn curve"
        whats = [w for _, w in hs if w]
        detail = (whats[0] + (" each" if len(hs) > 1 and not whats[0].endswith("only") else "")) if len(set(whats)) == 1 and whats else (_join([f"{h}: {w}" for h, w in hs if w]) if len(hs) > 1 else (hs[0][1] if hs[0][1] else ""))
        bits.append(f"{names} {verb} {evidence}" + (f" ({detail})" if detail else ""))
    if gaps:
        bits.append(f"{_join(gaps)} {'has' if len(gaps) == 1 else 'have'} no cell yet")
    s = "; ".join(bits)
    s = s[:1].upper() + s[1:] + "."
    if quals and len(set(quals)) == 1 and quals[0] != "live":
        s += " All of them are " + ("proposals, not in the runtime index yet." if quals[0] == "proposal" else "reused bindings.")
    none_any = [label_case(row["label"].split(":")[0]) for row in rows if all(row["cells"].get(h["slug"], {}).get("state") in (None, "gap", "no_curve", "not_named", "withheld", "exposure_modifier") for h in five)]
    if none_any:
        s += f" {len(none_any)} part{'s' if len(none_any) != 1 else ''} of the {noun} {'have' if len(none_any) != 1 else 'has'} no curve under any hazard: {_join(none_any)}."
    return s
