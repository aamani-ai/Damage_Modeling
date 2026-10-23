"""Sources as links, and the evidence of the older cells read from where it actually lives (iteration 2, brief §3.2, §3.5).

Every source carries `title`, `url`, `doi`, `locator`, and a resolved `link` with its kind: a web URL, a DOI, a file in
this repository at the build commit, a file in the Hazard repository, or none. A source with no link is shown with its
title and locator and the words *no link in the register* — a finding, not a blank.

The four oldest cells keep their sources, claims and assumptions in workbook sheets (`Sources`, `Evidence_Log`,
`*_Evidence_Params`, `Open_Seams*`, `*Assumption_Register`; standard 02 §5) and in the dossier's tables. This module
reads them and marks each item `extracted-from-workbook` or `from-dossier`; the tier table's source IDs are matched to
the extracted sources by their words, and every unmatched ID is reported, never dropped. A cell where neither works says
*not yet extracted*.
"""
from __future__ import annotations

import functools
import re
import subprocess
from pathlib import Path, PurePosixPath
from typing import Any

from .common import CHECKOUT, ROOT, rel, sha256

try:
    import openpyxl  # type: ignore
except ImportError:  # pragma: no cover
    openpyxl = None

HAZARD_REPO_URL = "https://github.com/aamani-ai/Hazard_Modeling"
DOI_RE = re.compile(r"\b(10\.\d{4,9}/[^\s\"'<>,;)\]]+)")
URL_RE = re.compile(r"https?://[^\s\"'<>|]+")
GENERIC_TOKENS = {"e", "src", "ev", "evid", "the", "of", "and", "for", "a", "an", "to", "in", "on", "with", "by", "or",
                  "public", "result", "results", "context", "declared", "method", "guidance", "baseline", "page", "paper",
                  "report", "memo", "sheet", "data", "v1", "v2", "v0", "doc", "docs", "local", "url", "pdf", "www", "com",
                  "gov", "org", "html", "http", "https", "en"}
NOT_YET_EXTRACTED = "Sources are in the dossier prose and the workbook; not yet extracted — open the dossier."
NO_LINK = "no link in the register"


# ─── links ────────────────────────────────────────────────────────────────────────────────────────────────────

def short_title(citation: str | None, fallback: str) -> str:
    """A reader's title: the citation's first sentence, or the whole citation when it is short."""
    c = re.sub(r";?\s*commit\s+[0-9a-f]{7,40}\b", "", (citation or "").strip()).strip(" ;")
    if not c:
        return fallback
    if len(c) <= 110:
        return c
    m = re.search(r"\.\s+(?=[A-Z])", c[20:])
    cut = c[: 20 + m.start()] if m else c[:110].rsplit(" ", 1)[0] + "…"
    return cut if len(cut) >= 24 else c[:110].rsplit(" ", 1)[0] + "…"


MARKER_RE = re.compile(r"^(NO_SOURCE|NO_DIRECT_SOURCE|GOVERNANCE_CONTRACT|SCOPE$|OPEN_VALUE)", re.I)


@functools.lru_cache(maxsize=1)
def _git_visible() -> frozenset[str] | None:
    """Repository files git can see (tracked, or new and not ignored), with their folders. A preview build reads the
    working tree, which may hold gitignored local mirrors (docs/source_drops/extracted/) and gitignored links to other
    checkouts (infrasure-damage-curves/ → the old repository); a main build reads an export that never has them."""
    if ROOT != CHECKOUT:
        return None
    done = subprocess.run(["git", "-C", str(CHECKOUT), "ls-files", "-z", "--cached", "--others", "--exclude-standard"], capture_output=True, text=True)
    if done.returncode != 0:
        return None
    files = [f for f in done.stdout.split("\0") if f]
    return frozenset(files) | frozenset(str(d) for f in files for d in PurePosixPath(f).parents if str(d) != ".")


def _inside(p: Path) -> bool:
    """The path stays inside the repository being read once every link on the way is followed."""
    try:
        p.resolve().relative_to(ROOT.resolve())
        return True
    except ValueError:
        return False


def _visible(p: Path) -> bool:
    """A path a link may point at: inside the repository and seen by git. A preview never links a gitignored mirror or a
    file reached through a gitignored link to another checkout; on GitHub neither exists."""
    if not _inside(p):
        return False
    seen = _git_visible()
    return seen is None or rel(p) in seen


def _repo_path(path: str, base_dir: Path | None) -> Path | None:
    """A repository file named by a register: an exact path, a path beside the register, or a unique file name under docs/.
    Only a path git sees inside this repository counts; a path through a folder outside it is no repository file."""
    p = path.strip("/").split("#")[0].split(";")[0].strip()
    if not p or p.startswith("http"):
        return None
    for cand in ([ROOT / p] + ([base_dir / p] if base_dir else [])):
        if cand.exists() and _visible(cand):
            return cand
    name = Path(p).name
    if "." in name and len(name) > 6:
        hits = [h for h in (ROOT / "docs").rglob(name) if h.is_file() and _visible(h)]
        for keep in (lambda h: "/archive/" not in str(h) and "/source_drops/" not in str(h), lambda h: "/source_drops/" not in str(h), lambda h: True):
            kept = [h for h in hits if keep(h)]
            if len(kept) == 1:
                return kept[0]
            if kept:
                break
    return None


def resolve_link(src: dict[str, Any], repo_url: str, commit: str, base_dir: Path | None = None) -> dict[str, Any]:
    """Fill `title`, `doi`, `link`, `link_kind` on a source from its url / citation / locator; never drops a field."""
    url = (src.get("url") or "").strip()
    citation = src.get("citation") or ""
    doi = None
    m = DOI_RE.search(url) or DOI_RE.search(citation)
    if m:
        doi = m.group(1).rstrip(".")
    link, kind = None, "none"
    if re.match(r"^https?://", url):
        link, kind = url, ("doi" if "doi.org/" in url else "url")
    elif url:
        path = url.strip("/").split("#")[0].split(";")[0].strip()
        if path.startswith("Hazard_modeling/"):
            # checked on Hazard's origin/main (G18): on main → a link at that commit; moved → the new place; else no link
            from . import hazard_main
            h = hazard_main.lookup(path[len("Hazard_modeling/"):])
            src["hazard"] = h
            link, kind = h["link"], {"on_main": "hazard_file", "moved": "hazard_moved", "on_branch": "hazard_branch", "unchecked": "hazard_file"}[h["state"]]
        else:
            rp = _repo_path(path, base_dir)
            if rp is not None:
                link = f"{repo_url}/{'tree' if rp.is_dir() else 'blob'}/{commit}/{rel(rp)}"
                kind = "repo_file"
                src["repo_path"] = rel(rp)
            else:
                found = URL_RE.search(url)
                if found:
                    link, kind = found.group(0), "url"
    if not link and doi:
        link, kind = f"https://doi.org/{doi}", "doi"
    if not link:
        found = URL_RE.search(citation)
        if found:
            link, kind = found.group(0), "url"
    if not link and src.get("locator"):
        rp = _repo_path(str(src["locator"]), base_dir)
        if rp is not None:
            link, kind = f"{repo_url}/blob/{commit}/{rel(rp)}", "repo_file"
            src["repo_path"] = rel(rp)
    src["doi"] = doi
    src["link"] = link
    src["link_kind"] = kind
    src.setdefault("title", short_title(citation, src.get("id", "")))
    if MARKER_RE.match(src.get("id") or ""):
        src["kind"] = "marker"  # a provenance marker ("no direct source"), not a document: nothing to link
    src["no_link_sentence"] = None if (link or src.get("kind") == "marker") else ((src.get("hazard") or {}).get("words") or NO_LINK)
    return src


def link_counts(sources: list[dict[str, Any]]) -> dict[str, int]:
    docs = [s for s in sources if s.get("kind") not in ("marker", "placeholder")]
    linked = sum(1 for s in docs if s.get("link"))
    locator_only = sum(1 for s in docs if not s.get("link") and s.get("locator"))
    neither = len(docs) - linked - locator_only
    return {"total": len(docs), "linked": linked, "locator_only": locator_only, "neither": neither,
            "markers": sum(1 for s in sources if s.get("kind") == "marker"), "placeholders": sum(1 for s in sources if s.get("kind") == "placeholder")}


# ─── workbook and dossier extraction ──────────────────────────────────────────────────────────────────────────

def _rows(ws: Any) -> list[list[str]]:
    out = []
    for r in ws.iter_rows(values_only=True):
        cells = ["" if v is None else str(v).strip() for v in r]
        while cells and cells[-1] == "":
            cells.pop()
        if any(cells):
            out.append(cells)
    return out


def _header_index(rows: list[list[str]], wanted: tuple[str, ...]) -> int | None:
    """The row that is the table header: it contains at least two of the wanted column names (lower-cased)."""
    for i, r in enumerate(rows[:8]):
        low = [c.lower().strip() for c in r]
        if sum(1 for w in wanted if any(w == c or (w in c and len(c) <= 40) for c in low)) >= 2:
            return i
    return None


def _table(rows: list[list[str]], wanted: tuple[str, ...]) -> list[dict[str, str]]:
    hi = _header_index(rows, wanted)
    if hi is None:
        return []
    header = [c.strip().lower().replace(" / ", "_").replace(" ", "_").replace("-", "_") for c in rows[hi]]
    out = []
    for r in rows[hi + 1:]:
        if not any(r):
            continue
        row = {header[i]: (r[i] if i < len(r) else "") for i in range(len(header))}
        row["_extra"] = " ".join(r[len(header):]) if len(r) > len(header) else ""
        out.append(row)
    return out


def _first(row: dict[str, str], *keys: str) -> str:
    for k in keys:
        for col, v in row.items():
            if col == k and v:
                return v
    for k in keys:
        for col, v in row.items():
            if k in col and v:
                return v
    return ""


def _slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_").upper()[:60]


def _url_in(*texts: str) -> str | None:
    for t in texts:
        m = URL_RE.search(t or "")
        if m:
            return m.group(0).rstrip(".,;)")
    return None


def extract_workbook(wb_path: Path, cell_token: str) -> dict[str, Any]:
    """Sources, claims, assumptions and open seams from the workbook's sheets (standard 02 §5)."""
    out: dict[str, Any] = {"sources": [], "claims": [], "assumptions": [], "open_seams": [], "sheets_read": []}
    if openpyxl is None or not wb_path.exists():
        return out
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    where = rel(wb_path)
    by_url: dict[str, dict[str, Any]] = {}

    def add_source(sid: str, title: str, url: str | None, role: str, notes: str, sheet: str, extra: dict[str, Any] | None = None) -> dict[str, Any]:
        key = (url or "").rstrip("/").lower() or f"id:{sid.lower()}"
        if key in by_url:
            s = by_url[key]
            s["aliases"] = sorted(set(s.get("aliases", [])) | {sid})
            if not s.get("role") and role:
                s["role"] = role
            return s
        s = {"id": sid, "title": title or sid, "citation": title or sid, "url": url, "locator": None, "role": role or None, "tier": None,
             "permitted": None, "prohibited": notes or None, "decision": None, "accessed_on": None,
             "provenance": "extracted-from-workbook", "from": {"file": where, "sheet": sheet}, "aliases": [sid]}
        for k, v in (extra or {}).items():
            if v is not None or k not in s:
                s[k] = v
        by_url[key] = s
        out["sources"].append(s)
        return s

    for name in wb.sheetnames:
        low = name.lower()
        ws = wb[name]
        if low == "sources":
            rows = _rows(ws)
            table = _table(rows, ("source_id", "title", "url", "publisher", "location_or_url"))
            for r in table:
                sid = _first(r, "source_id")
                url = _url_in(_first(r, "url", "location_or_url"))
                title = _first(r, "title") or ""
                stype = _first(r, "source_type")
                shifted = not url and _url_in(stype, r.get("_extra", ""))
                if shifted or not sid:
                    # a shifted row (the hail workbook): the name sits in the first column, the URL in the second
                    first = next((v for v in r.values() if v), "")
                    url = url or _url_in(*r.values())
                    title = first if not first.startswith("http") else (title or sid)
                    sid = _slug(first) if first and not first.startswith("http") else (sid or _slug(url or ""))
                    stype = "public_url" if url else stype
                if not sid:
                    continue
                extra: dict[str, Any] = {"locator": None if url else (_first(r, "location_or_url") or None), "source_type": stype or None}
                if "placeholder" in stype.lower():
                    extra["kind"] = "placeholder"  # an input still to be supplied, not a document behind a number
                elif not url:
                    extra["url"] = _first(r, "url", "location_or_url") or None  # an internal document named by path; resolved to a repository file if it exists
                add_source(sid, title, url, _first(r, "used_for", "role", "how_used"), _first(r, "notes"), name, extra)
            out["sheets_read"].append(name)
        elif low == "evidence_log":
            rows = _rows(ws)
            for r in _table(rows, ("evidence_id", "source_name", "source_location_or_url")):
                applies = _first(r, "applies_to")
                if cell_token and applies and cell_token.lower() not in applies.lower() and "all" not in applies.lower():
                    continue
                url = _url_in(_first(r, "source_location_or_url"))
                if not url:
                    continue  # foundation docs and placeholders are not sources for this cell's numbers
                s = add_source(_first(r, "evidence_id"), _first(r, "source_name"), url, _first(r, "source_role"), _first(r, "caveat"), name)
                out["claims"].append({"id": _first(r, "evidence_id"), "text": f"{_first(r, 'source_role')}: {_first(r, 'parameter_or_field')}".strip(": "),
                                      "type": _first(r, "evidence_class"), "source_ids": [s["id"]], "locator": None, "tier": None,
                                      "strength": _first(r, "trust_tier"), "parameter_or_rule": _first(r, "parameter_or_field"), "status": None,
                                      "update_trigger": None, "caveat": _first(r, "caveat"), "provenance": "extracted-from-workbook", "from": {"file": where, "sheet": name}})
            out["sheets_read"].append(name)
        elif "evidence_params" in low or low == "evidence_params":
            rows = _rows(ws)
            table = _table(rows, ("evidence_id", "source", "parameter_or_rule", "supports", "evidence_item", "parameter_id"))
            for i, r in enumerate(table, start=1):
                url = _url_in(_first(r, "source_url"))
                src_name = _first(r, "source", "evidence_source", "source_id")
                ids: list[str] = []
                if url:
                    ids = [add_source(_first(r, "evidence_id") or _slug(src_name), src_name, url, _first(r, "source_class"), "", name)["id"]]
                elif _first(r, "source_id"):
                    ids = [_first(r, "source_id")]
                elif _first(r, "evidence_id") and ";" in _first(r, "evidence_id"):
                    ids = [x.strip() for x in _first(r, "evidence_id").split(";") if x.strip()]
                elif _first(r, "evidence_id") and not _first(r, "evidence_item") and not _first(r, "quoted_or_extracted_fact"):
                    ids = [_first(r, "evidence_id")]  # the flood workbook: evidence_id is the source id list
                text = _first(r, "quoted_or_extracted_fact", "evidence_item", "supports", "parameter_or_rule")
                cid = _first(r, "evidence_id", "parameter_id") or f"{name}-{i}"
                if ids == [cid] and _first(r, "parameter_id"):
                    cid = _first(r, "parameter_id")
                out["claims"].append({"id": cid, "text": text, "type": _first(r, "source_class", "adopted_role"), "source_ids": ids or ([_slug(src_name)] if src_name else []),
                                      "source_name": src_name or None, "locator": None, "tier": None, "strength": _first(r, "trust_tier", "evidence_strength", "curve_changing"),
                                      "parameter_or_rule": _first(r, "parameter_or_rule", "parameter_use", "parameter_rule_affected"), "how_used": _first(r, "how_used", "supports", "source_interpretation"),
                                      "numeric": _first(r, "numeric_extract"), "status": None, "update_trigger": None,
                                      "provenance": "extracted-from-workbook", "from": {"file": where, "sheet": name}})
            out["sheets_read"].append(name)
        elif "assumption_register" in low:
            rows = _rows(ws)
            for r in _table(rows, ("assumption_id", "assumption")):
                out["assumptions"].append({"id": _first(r, "assumption_id"), "text": _first(r, "assumption"), "status": _first(r, "status", "strength", "materiality"),
                                           "confidence": _first(r, "confidence"), "trigger": _first(r, "replacement_trigger", "update_trigger", "update trigger"),
                                           "effect": _first(r, "effect_if_changed", "effect_if_revised", "affects"), "from": {"file": where, "sheet": name}})
            out["sheets_read"].append(name)
        elif low.startswith("open_seams"):
            rows = _rows(ws)
            for r in _table(rows, ("seam_id", "open_seam_id", "description", "seam_name")):
                cid = _first(r, "cell_id")
                if cell_token and cid and cell_token.lower() != cid.lower():
                    continue
                out["open_seams"].append({"id": _first(r, "open_seam_id", "seam_id"), "name": _first(r, "seam_name") or _first(r, "description")[:80],
                                          "description": _first(r, "description"), "severity": _first(r, "severity", "materiality"), "status": _first(r, "status"),
                                          "resolution": _first(r, "proposed_resolution", "resolution_path"), "from": {"file": where, "sheet": name}})
            out["sheets_read"].append(name)
    return out


def extract_dossier_tables(path: Path) -> dict[str, Any]:
    """Any pipe table in the dossier whose rows carry a URL: a source per row, with the table's heading as the section."""
    out: dict[str, Any] = {"sources": [], "claims": [], "sections": []}
    if not path.exists():
        return out
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    heading = ""
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("#"):
            heading = ln.lstrip("#").strip()
        if ln.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in ln.strip().strip("|").split("|")]
            j = i + 2
            rows = []
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            url_col = next((k for k, h in enumerate(header) if re.search(r"url|link", h, re.I)), None)
            if url_col is not None and any(len(r) > url_col and URL_RE.search(r[url_col]) for r in rows):
                out["sections"].append(heading)
                for r in rows:
                    if len(r) <= url_col:
                        continue
                    url = _url_in(r[url_col])
                    name_col = next((k for k, h in enumerate(header) if re.search(r"source|evidence|name", h, re.I) and k != url_col), 0)
                    sid_col = next((k for k, h in enumerate(header) if re.search(r"\bid\b", h, re.I)), None)
                    name = r[name_col].strip("`") if name_col < len(r) else ""
                    sid = (r[sid_col].strip("`") if sid_col is not None and sid_col < len(r) else "") or _slug(name)
                    supports = next((r[k] for k, h in enumerate(header) if re.search(r"supports|what it", h, re.I) and k < len(r) and not re.search(r"not", h, re.I)), "")
                    limits = next((r[k] for k, h in enumerate(header) if re.search(r"does \*\*not\*\*|not support", h, re.I) and k < len(r)), "")
                    out["sources"].append({"id": sid, "title": name or sid, "citation": name or sid, "url": url, "locator": None if url else (r[url_col] or None),
                                           "role": None, "tier": None, "permitted": supports or None, "prohibited": limits or None, "decision": None, "accessed_on": None,
                                           "provenance": "from-dossier", "from": {"file": rel(path), "section": heading}, "aliases": [sid]})
                    if supports:
                        out["claims"].append({"id": f"{sid}·dossier", "text": supports, "type": "dossier evidence map", "source_ids": [sid], "locator": heading, "tier": None,
                                              "strength": None, "parameter_or_rule": None, "status": None, "update_trigger": None,
                                              "provenance": "from-dossier", "from": {"file": rel(path), "section": heading}})
            i = j
            continue
        i += 1
    return out


# ─── matching the tier table's source IDs to the extracted sources ───────────────────────────────────────────

def _tokens(text: str) -> set[str]:
    toks = set()
    for t in re.split(r"[^A-Za-z0-9]+", (text or "").lower()):
        if not t or t in GENERIC_TOKENS:
            continue
        t = t[:-1] if len(t) > 4 and t.endswith("s") else t
        toks.add(t)
    return toks


def match_ids(ids: list[str], sources: list[dict[str, Any]], claims: list[dict[str, Any]]) -> tuple[dict[str, str], list[str]]:
    """Each tier-table source ID → the extracted source whose words it shares most; unmatched IDs are returned, not guessed."""
    pools: list[tuple[dict[str, Any], set[str], set[str]]] = []
    for s in sources:
        own = _tokens(" ".join([s.get("id", ""), s.get("title", ""), *s.get("aliases", [])]))
        pool = own | _tokens(" ".join([s.get("url", "") or "", s.get("role") or ""]))
        for c in claims:
            if s["id"] in c.get("source_ids", []) or (c.get("source_name") and c["source_name"] == s.get("title")):
                pool |= _tokens(" ".join([c.get("id", ""), c.get("text", ""), c.get("parameter_or_rule") or "", c.get("how_used") or ""]))
        pools.append((s, own, pool))
    weight_of = lambda toks: sum(1.0 if len(t) <= 3 else 2.0 if len(t) <= 5 else 3.0 for t in toks)  # noqa: E731 — rarer, longer words weigh more than "wind" or "pv"
    matched: dict[str, str] = {}
    unmatched: list[str] = []
    for sid in ids:
        exact = next((s for s in sources if s["id"] == sid or sid in s.get("aliases", [])), None)
        if exact:
            matched[sid] = exact["id"]
            continue
        want = _tokens(sid)
        if not want:
            unmatched.append(sid)
            continue
        scored = []
        for order, (s, own, pool) in enumerate(pools):
            common = want & pool
            if not common:
                continue
            if len(common) < 2 and not any(len(t) >= 5 and t in own for t in common):
                continue  # one shared word is a match only when it is distinctive and in the source's own name
            score = weight_of(common) / max(1.0, weight_of(want))
            scored.append((score, weight_of(want & own), 1 if s.get("link") else 0, -order, s))
        scored.sort(key=lambda x: x[:4], reverse=True)
        if scored and scored[0][0] >= 0.45:
            matched[sid] = scored[0][4]["id"]
        else:
            unmatched.append(sid)
    return matched, unmatched


def repo_file_sources(ids: list[str], repo_url: str, commit: str) -> list[dict[str, Any]]:
    """Source IDs that name a file in this repository (e.g. SOLAR_WIND_VALUE_BREAKDOWN → docs/method/value_basis/…xlsx)."""
    out = []
    candidates = list((ROOT / "docs" / "method" / "value_basis").glob("*")) if (ROOT / "docs" / "method" / "value_basis").is_dir() else []
    for sid in ids:
        want = _tokens(sid)
        for p in candidates:
            if want and want <= _tokens(p.stem) | {"xlsx"}:
                out.append({"id": sid, "title": p.name, "citation": f"{p.name} (this repository)", "url": rel(p), "locator": rel(p), "role": "value basis",
                            "tier": None, "permitted": None, "prohibited": None, "decision": None, "accessed_on": None,
                            "provenance": "repository-file", "from": {"file": rel(p)}, "aliases": [sid]})
                break
    return out


def extract_for_entry(entry: Any, registers: dict[str, Any], repo_url: str, commit: str) -> dict[str, Any] | None:
    """For a cell with parameters but no source register: read the workbook and the dossier, match the tier table's
    source IDs, and return the extracted registers with their provenance — or None when there is nothing to read."""
    if registers.get("sources") or entry.kind != "cell" or not entry.folder:
        return None
    books = sorted(entry.folder.glob("*.xlsx"))
    a = entry.artifact or {}
    cand = a.get("source_dossier")
    dossier = ROOT / cand if cand and (ROOT / cand).exists() else next((p for p in sorted(entry.folder.glob("*derivation*.md"))), None)
    if not books and not dossier:
        return None
    cell_token = (a.get("cell_id") or entry.id).replace("-", "_")
    wb = extract_workbook(books[-1], cell_token) if books else {"sources": [], "claims": [], "assumptions": [], "open_seams": [], "sheets_read": []}
    ds = extract_dossier_tables(dossier) if dossier else {"sources": [], "claims": [], "sections": []}
    # one list: workbook first, then dossier rows whose URL the workbook does not already carry
    seen_urls = {(s.get("url") or "").rstrip("/").lower() for s in wb["sources"] if s.get("url") and str(s["url"]).startswith("http")}
    seen_ids = {s["id"].lower() for s in wb["sources"]}
    sources = list(wb["sources"])
    for s in ds["sources"]:
        key = (s.get("url") or "").rstrip("/").lower()
        if (key and key in seen_urls) or s["id"].lower() in seen_ids:
            continue
        seen_urls.add(key)
        seen_ids.add(s["id"].lower())
        sources.append(s)
    claims = list(wb["claims"]) + list(ds["claims"])
    ids = sorted({sid for p in registers.get("parameters") or [] for sid in p.get("source_ids", [])})
    sources += repo_file_sources([i for i in ids], repo_url, commit)
    for s in sources:
        resolve_link(s, repo_url, commit, entry.folder)
    matched, unmatched = match_ids(ids, [s for s in sources if s.get("kind") not in ("placeholder",)], claims)
    if not sources and not claims:
        return {"status": "not_yet_extracted", "sentence": NOT_YET_EXTRACTED, "dossier": rel(dossier) if dossier else None,
                "workbook": rel(books[-1]) if books else None, "sources": [], "claims": [], "assumptions": [], "open_seams": [],
                "matched": {}, "unmatched": ids, "sheets_read": []}
    return {"status": "extracted", "sentence": None, "dossier": rel(dossier) if dossier else None, "workbook": rel(books[-1]) if books else None,
            "sources": sources, "claims": claims, "assumptions": wb["assumptions"], "open_seams": wb["open_seams"],
            "matched": matched, "unmatched": unmatched, "sheets_read": wb["sheets_read"], "dossier_sections": ds["sections"],
            "files": [{"path": rel(p), "sha256": sha256(p)} for p in ([books[-1]] if books else []) + ([dossier] if dossier else [])]}
