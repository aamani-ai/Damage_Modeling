"""Validator for the noncanonical convective wind × real estate model-v0.1 / docs-r1 package.

Checks: required files; JSON schemas (bundle v3, capability v3, damage emit v2, when jsonschema is installed);
embedded vs standalone capability; rectangular LF registers and resolvable source IDs; the EF table (bounds ordered,
costs on the Hazus ladder and non-decreasing); the ordinates re-derived from the EF table; artifact records vs the
governed ordinates (knots, range, monotone, [0, 1], straight-line zero at and below the onset); loader compatibility
(one selector, unique routing per pathway, one range per pathway, distinct axis fields); builder reproducibility;
every known-answer test through the reference evaluator; non-promotion; the profile binding; workbook sheets; links.

Usage (from the repository root):
    python scripts/reference_helpers/validate_convective_wind_real_estate_v0_1_package.py
Exit code 0 when every check passes.
"""
from __future__ import annotations

import csv
import json
import math
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from convective_wind_real_estate_curve_eval import DamageCallError, evaluate_damage_call  # noqa: E402
import build_convective_wind_real_estate_v0_1_package as builder  # noqa: E402
import derive_convective_wind_real_estate_ef_curves as derive  # noqa: E402

CELL_ID = builder.CELL_ID
CELL = ROOT / f"docs/cells/{CELL_ID}"
PROP = CELL / "proposed"
TAG = builder.TAG
SCHEMAS = ROOT / "docs/contracts/schemas"
HAZUS_LADDER = {0.01, 0.02, 0.05, 0.10, 0.50, 1.00}
results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((name, bool(ok), detail))


def load(path: Path):
    with open(path) as fh:
        return json.load(fh)


def csv_rows(path: Path) -> list[list[str]]:
    with open(path, newline="") as fh:
        return list(csv.reader(fh))


def dict_rows(path: Path) -> list[dict[str, str]]:
    with open(path, newline="") as fh:
        return list(csv.DictReader(fh))


def main() -> int:
    files = {
        "artifact": PROP / f"{TAG}__curve_artifact.json",
        "capability": PROP / f"{TAG}__capability.json",
        "kats": PROP / f"known_answer_tests_{TAG}.json",
        "workbook": PROP / f"damage_curve_records_{TAG}.xlsx",
        "dossier": PROP / f"{CELL_ID}_curve_derivation_dossier__model_v0_1__docs_r1.md",
        "metadata_spec": PROP / f"{CELL_ID}_damage_code_metadata_spec__model_v0_1__docs_r1.md",
        "decision_log": PROP / f"DECISION_LOG_{TAG}.md",
        "validation_report": PROP / f"VALIDATION_REPORT_{TAG}.md",
        "return_package": PROP / f"RETURN_PACKAGE_{TAG}.md",
        "source_register": PROP / f"SOURCE_REGISTER_{TAG}.csv",
        "claim_register": PROP / f"CLAIM_PARAMETER_REGISTER_{TAG}.csv",
        "tiers": PROP / f"PARAMETER_TIER_TABLE_{TAG}.csv",
        "crosswalk": PROP / f"VALUE_CROSSWALK_{TAG}.csv",
        "ef_table": PROP / f"EF_DEGREE_OF_DAMAGE_TABLE_{TAG}.csv",
        "ordinates": PROP / f"CURVE_ORDINATES_{TAG}.csv",
        "sensitivity": PROP / f"SENSITIVITY_{TAG}.csv",
        "nri_check": PROP / f"NRI_CHECK_{TAG}.csv",
        "nri_receipt": PROP / f"NRI_CHECK_RECEIPT_{TAG}.json",
        "binding": PROP / f"PROFILE_BINDING_{TAG}.csv",
    }
    missing = [k for k, p in files.items() if not p.exists()]
    check("required files exist", not missing, ", ".join(missing))
    if missing:
        return report()

    art, cap, kats = load(files["artifact"]), load(files["capability"]), load(files["kats"])
    pathways = {p["pathway_id"]: p for p in art["pathways"]}
    recs = [r for p in art["pathways"] for r in p["curve_records"]]
    check("embedded capability equals standalone capability", art["capability_declaration"] == cap)
    for key in ("source_dossier", "source_workbook", "known_answer_tests", "source_register",
                "claim_parameter_register", "value_crosswalk"):
        check(f"artifact path resolves: {key}", (ROOT / art[key]).exists(), art[key])

    # schemas
    try:
        from jsonschema import Draft202012Validator
        from referencing import Registry, Resource
        sch = {n: load(SCHEMAS / f"{n}.schema.json") for n in
               ("curve_artifact_bundle.v3", "capability_declaration.v3", "damage_emit.v2")}
        reg = Registry().with_resources([(s["$id"], Resource.from_contents(s)) for s in sch.values()])
        validators = {n: Draft202012Validator(s, registry=reg) for n, s in sch.items()}
        errs = list(validators["curve_artifact_bundle.v3"].iter_errors(art))
        check("artifact validates against bundle v3", not errs, "; ".join(e.message[:120] for e in errs[:3]))
        errs = list(validators["capability_declaration.v3"].iter_errors(cap))
        check("capability validates against capability v3", not errs, "; ".join(e.message[:120] for e in errs[:3]))
        emit_validator = validators["damage_emit.v2"]
    except ModuleNotFoundError:
        check("jsonschema available for schema checks", False, "install jsonschema>=4.17")
        emit_validator = None

    # registers
    for k in ("source_register", "claim_register", "tiers", "crosswalk", "ef_table", "ordinates", "sensitivity",
              "nri_check", "binding"):
        widths = {len(r) for r in csv_rows(files[k])}
        check(f"register rectangular: {files[k].name}", len(widths) == 1, str(widths))
        check(f"register uses LF line endings: {files[k].name}", b"\r" not in files[k].read_bytes())
    sources = {r[0] for r in csv_rows(files["source_register"])[1:]}
    refs = {ref.split("#")[0] for r in recs for ref in r["source_parameter_refs"]}
    check("curve source refs resolve to the source register", refs <= sources, str(refs - sources))
    for k, col in (("claim_register", "source_ids"), ("tiers", "source_ids"), ("ef_table", "source_ids"),
                   ("ordinates", "source_ids")):
        ids = {s for r in dict_rows(files[k]) for s in r[col].split(";") if s}
        check(f"{files[k].name.split('_convective')[0]} source IDs resolve", ids <= sources, str(ids - sources))

    # EF degree-of-damage table
    ef = dict_rows(files["ef_table"])
    bad = [f"{r['reference_building_id']} DOD{r['dod']}" for r in ef
           if not float(r["lower_bound_mph"]) < float(r["expected_mph"]) < float(r["upper_bound_mph"])
           or float(r["cost_ratio"]) not in HAZUS_LADDER]
    check("EF table: LB < EXP < UB and every cost on the Hazus ladder (1/2/5/10/50/100%)", not bad, str(bad))
    table = derive.dod_table()
    nondec = all(float(a["cost_ratio"]) <= float(b["cost_ratio"]) for rows in table.values() for a, b in zip(rows, rows[1:]))
    check("EF table: costs non-decreasing with degree of damage", nondec)
    check("EF table: degree counts FR12 10, MBS 8, LRB 7",
          {b: len(r) for b, r in table.items()} == {"RE_RES_SFD_1S_SLAB": 10, "RE_IND_WAREHOUSE": 8, "RE_COM_OFFICE_3S_SLAB": 7})

    # ordinates re-derived from the EF table
    def as_text(data: list[list]) -> list[list[str]]:
        return [[str(c) for c in row] for row in data]

    check("ordinates CSV equals a fresh derivation from the EF table",
          as_text(derive.ordinates(table)) == csv_rows(files["ordinates"]))
    check("sensitivity CSV equals a fresh derivation", as_text(derive.sensitivity(table)) == csv_rows(files["sensitivity"]))

    # records
    pts = builder.ordinates()
    knots = [float(v) for v in range(0, 256)]
    mism = [r["curve_id"] for r in recs
            if r["parameters"]["points"] != pts[(r["selector_match"]["reference_building_id"], r["pathway_id"])]]
    check(f"artifact records equal the governed ordinates ({builder.EXPECTED_RECORDS} records)",
          not mism and len(recs) == builder.EXPECTED_RECORDS, str(mism))
    check("every record has 1-mph knots 0..255", all([p[0] for p in r["parameters"]["points"]] == knots for r in recs))
    check("every value is within [0, 1]", all(0.0 <= p[1] <= 1.0 for r in recs for p in r["parameters"]["points"]))
    check("every record is non-decreasing", all(b[1] >= a[1] for r in recs for a, b in zip(r["parameters"]["points"], r["parameters"]["points"][1:])))
    sl = [r for r in recs if r["pathway_id"] == builder.SL]
    check("straight-line records are 0 at and below the 96 mph onset and above 0 just after",
          all(all(y == 0.0 for x, y in r["parameters"]["points"] if x <= builder.ONSET) and r["parameters"]["points"][97][1] > 0 for r in sl))
    tor = {r["selector_match"]["reference_building_id"]: r for r in recs if r["pathway_id"] == builder.TOR}
    check("tornado record >= straight-line record at every knot (onset removes damage only)",
          all(t[1] >= s[1] for r in sl for s, t in zip(r["parameters"]["points"], tor[r["selector_match"]["reference_building_id"]]["parameters"]["points"])))
    check("each record carries its pathway flags", all(set(builder.PATHWAY_FLAGS[r["pathway_id"]]) <= set(r["metadata_flags"]) for r in recs))

    # loader compatibility (Hazard generic v3 loader)
    for pid, pw in pathways.items():
        prs = pw["curve_records"]
        check(f"{pid}: routing field is exactly the one required selector",
              pw["record_routing_fields"] == [s["field"] for s in pw["selector_logic"] if s["required"] is True] == ["reference_building_id"])
        combos = [r["selector_match"]["reference_building_id"] for r in prs]
        check(f"{pid}: every building routes to exactly one record", sorted(combos) == sorted(builder.BUILDINGS))
        check(f"{pid}: all records share the pathway valid_range", all(r["valid_range"] == pw["hazard_axis"]["valid_range"] for r in prs))
        check(f"{pid}: records use the pathway axis field", all(r["x_axis"] == pw["hazard_axis"]["preferred_input_field"] for r in prs))
    check("the two pathways read different axis fields",
          pathways[builder.SL]["hazard_axis"]["preferred_input_field"] != pathways[builder.TOR]["hazard_axis"]["preferred_input_field"])

    # reproducibility
    check("builder reproduces the artifact JSON", builder.artifact(pts) == art)
    check("builder reproduces the known-answer tests", builder.known_answer_tests(pts) == kats)

    # known-answer tests
    tol = kats["absolute_tolerance"]
    fails, n = [], 0
    for t in kats["formula_known_answer_tests"] + kats["withheld_tests"]:
        n += 1
        try:
            out = evaluate_damage_call(art, t["input"])
        except DamageCallError as e:
            fails.append(f"{t['test_id']}: error {e.code}")
            continue
        if emit_validator is not None:
            emit_errs = list(emit_validator.iter_errors(out))
            if emit_errs:
                fails.append(f"{t['test_id']}: emit schema {emit_errs[0].message[:80]}")
        res, exp = out["failure_unit_results"][0], t["expected"]
        if res["status"] != exp["status"]:
            fails.append(f"{t['test_id']}: status {res['status']}")
            continue
        if exp["status"] == "withheld":
            if exp["reason_code"] not in res["withheld_reason_codes"] or res["scalar_central_dr"] is not None:
                fails.append(f"{t['test_id']}: reasons {res['withheld_reason_codes']}")
            continue
        if res["curve_id"] != exp["curve_id"] or not math.isclose(res["scalar_central_dr"], exp["failure_unit_damage_ratio"], rel_tol=0, abs_tol=tol):
            fails.append(f"{t['test_id']}: {res['curve_id']} {res['scalar_central_dr']}")
        for f in exp.get("metadata_flags_include", []):
            if f not in res["metadata_flags"]:
                fails.append(f"{t['test_id']}: missing flag {f}")
    for t in kats["error_tests"]:
        n += 1
        try:
            evaluate_damage_call(art, t["input"])
            fails.append(f"{t['test_id']}: no error raised")
        except DamageCallError as e:
            if e.code != t["expected_error_code"]:
                fails.append(f"{t['test_id']}: {e.code}")
    check(f"known-answer tests pass ({n})", not fails, "; ".join(fails[:5]))

    # non-promotion
    index = (ROOT / "docs/contracts/machine_readable_artifact_index.json").read_text()
    registry = (ROOT / "docs/cells/VERSION_REGISTRY.md").read_text()
    check("proposal absent from the machine-readable artifact index", f"{TAG}__curve_artifact.json" not in index)
    check("proposal pin absent from the version registry", f"{CELL_ID}@model_v0_1__docs_r1" not in registry)
    check("proposal is not installed as current", not (CELL / "current" / f"{TAG}__curve_artifact.json").exists())
    check("artifact is noncanonical and proposed",
          art["canonical_runtime_artifact"] is False and art["schema_status"] == "proposed_draft"
          and art["promotion_status"] == "proposed" and cap["promotion_gate"]["status"] == "blocked")

    # NRI check receipt matches the ordinates used
    receipt = load(files["nri_receipt"])
    import hashlib
    check("NRI check receipt pins the current ordinates",
          receipt["inputs"]["ordinates"]["sha256"] == hashlib.sha256(files["ordinates"].read_bytes()).hexdigest())

    # profile binding
    by_id = {r["curve_id"]: r for r in recs}
    ids = sources | {r[0] for r in csv_rows(files["claim_register"])[1:]}
    binding = dict_rows(files["binding"])
    bad = []
    for b in binding:
        tag = f"{b['profile_id']}/{b['pathway_id']}"
        rec_ids = [c for c in b["record_ids"].split(";") if c]
        if b["applicability"] == "unavailable":
            if rec_ids:
                bad.append(f"{tag}: unavailable row names records")
            continue
        if len(rec_ids) != 1 or rec_ids[0] not in by_id or by_id[rec_ids[0]]["pathway_id"] != b["pathway_id"] \
                or by_id[rec_ids[0]]["selector_match"]["reference_building_id"] != b["damage_selector_reference_building_id"]:
            bad.append(f"{tag}: record missing or mismatched")
        bad += [f"{tag}: evidence {e} unknown" for e in b["evidence_ids"].split(";") if e and e not in ids]
    check("profile binding records and evidence resolve", binding and not bad, "; ".join(bad[:5]))
    check("profile binding covers every profile × pathway",
          {(b["profile_id"], b["pathway_id"]) for b in binding if b["record_ids"]}
          == {(s["profile"], p) for s in builder.BUILDINGS.values() for p in (builder.SL, builder.TOR)})
    check("profile binding pinned to one profile revision", len({(b["profile_revision"], b["profile_commit"]) for b in binding}) == 1)

    # workbook
    try:
        from openpyxl import load_workbook
        wb = load_workbook(files["workbook"], read_only=True)
        need = {"README", "curve_records", "reference_buildings", "ef_degrees_of_damage", "sources", "claims",
                "parameter_tiers", "value_crosswalk", "sensitivity", "nri_check", "known_answer_tests"}
        check("workbook has the expected sheets", need <= set(wb.sheetnames), str(need - set(wb.sheetnames)))
    except ModuleNotFoundError:
        check("openpyxl available for workbook check", False, "install openpyxl")

    # links
    broken = []
    for md in CELL.rglob("*.md"):
        for m in re.finditer(r"\]\(([^)#]+)(#[^)]*)?\)", md.read_text()):
            target = m.group(1)
            if not target.startswith("http") and not (md.parent / target).exists():
                broken.append(f"{md.name} -> {target}")
    check("Markdown links in the cell resolve", not broken, "; ".join(broken[:5]))
    return report()


def report() -> int:
    width = max(len(n) for n, _, _ in results)
    for name, ok, detail in results:
        print(f"{'PASS' if ok else 'FAIL'}  {name.ljust(width)}  {detail if not ok else ''}")
    failed = sum(1 for _, ok, _ in results if not ok)
    print(f"\n{len(results) - failed}/{len(results)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
