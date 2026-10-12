#!/usr/bin/env python3.12
"""Validate the noncanonical winter_storm_wind model v0.1 / docs r1 fail-closed scaffold."""

from __future__ import annotations

import csv
import hashlib
import json
import math
from pathlib import Path

import jsonschema
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
P = ROOT / "docs/cells/winter_storm_wind/proposed"
TAG = "model_v0_1__docs_r1"
ART = P / f"winter_storm_wind__{TAG}__curve_artifact.json"
CAP = P / f"winter_storm_wind__{TAG}__capability.json"
KAT = P / f"known_answer_tests_winter_storm_wind__{TAG}.json"
SRC = P / f"SOURCE_REGISTER_winter_storm_wind__{TAG}.csv"
CLM = P / f"CLAIM_PARAMETER_REGISTER_winter_storm_wind__{TAG}.csv"
PAR = P / f"PARAMETER_TIER_TABLE_winter_storm_wind__{TAG}.csv"
VAL = P / f"VALUE_CROSSWALK_winter_storm_wind__{TAG}.csv"
WB = P / f"damage_curve_records_winter_storm_wind__{TAG}.xlsx"
DOCS = ["README", "CHANGE_CLASSIFICATION", "DECISION_LOG", "BOUNDED_EVIDENCE_SEARCH_LOG", "LEGACY_EVIDENCE_INGESTION",
        "PRESSURE_TEST", "SEVEN_STEP_AUDIT", "PROMOTION_GATE_MATRIX", "VALIDATION_REPORT"]
METRICS = {"failure_unit_scalar_dr", "scenario_loss_given_value_basis", "scalar_eal", "pml", "var", "tvar"}
TIERS = {"T1_claims_or_field_calibrated", "T2_public_lab_standard_or_physics",
         "T3_engineering_proxy_or_adjacent_empirical", "T4_placeholder_or_expert_judgment"}
UNITS = {"WT_BLADE_ASSEMBLY", "WT_NACELLE_AND_HUB", "WT_TOWER_AND_EXTERNAL_FIXTURES", "WT_PAD_ELECTRICAL",
         "WT_COLLECTION_NETWORK", "WT_GSU_SUBSTATION", "WT_CONTROL_AND_MET_STATION", "WT_FOUNDATION", "WT_CIVIL_INFRA",
         "SUPPORT_FIELDWORK", "SUPPORT_TRANSPORT_LOGISTICS"}
DEFERRED = {"ice_with_concurrent_wind", "ice_shedding_impact", "snow_gravity_load_on_buildings", "extreme_cold_physical_effects"}
CHECKS = 0


def require(condition: bool, message: str) -> None:
    global CHECKS
    CHECKS += 1
    if not condition:
        raise AssertionError(message)


def rows(path: Path) -> list[dict]:
    require(b"\r\n" not in path.read_bytes(), f"{path.name}: CRLF")
    with path.open(newline="") as handle:
        out = list(csv.DictReader(handle))
    require(bool(out), f"{path.name}: empty")
    for row in out:
        require(None not in row and all(v is not None for v in row.values()), f"{path.name}: ragged row")
    return out


def scaffold_response(payload: dict) -> dict:
    """Reference fail-closed behaviour: contract checks first, then withhold every metric."""
    pathway = payload.get("pathway_id")
    if pathway is None:
        return {"error": "PATHWAY_ID_REQUIRED"}
    if pathway in DEFERRED:
        return {"error": "PATHWAY_DEFERRED_NOT_IMPLEMENTED"}
    if pathway != "ice_accretion":
        return {"error": "PATHWAY_ID_UNKNOWN"}
    if not payload.get("event_family_id"):
        return {"error": "EVENT_FAMILY_ID_REQUIRED"}
    if payload.get("failure_unit_id") not in UNITS:
        return {"error": "FAILURE_UNIT_ID_UNKNOWN"}
    return {m: {"value": None, "status": "withheld", "reason_codes": ["NO_RUNTIME_CURVE"]} for m in METRICS}


def main() -> int:
    artifact = json.loads(ART.read_text())
    capability = json.loads(CAP.read_text())
    schemas = ROOT / "docs/contracts/schemas"
    jsonschema.validate(artifact, json.loads((schemas / "curve_artifact_bundle.schema.json").read_text()))
    jsonschema.validate(capability, json.loads((schemas / "capability_declaration.schema.json").read_text()))
    require(artifact["cell_id"] == "winter_storm_wind", "cell id")
    require(artifact["lifecycle_state"] == "scaffold" and artifact["promotion_status"] == "proposed", "lifecycle")
    require(artifact["canonical_runtime_artifact"] is False, "became canonical")
    require(artifact["curve_records"] == [] and artifact["parameter_tier_table"] == [], "runtime curve leaked")
    require(artifact["ordinate_status"] == "withheld" and artifact["withheld_reason_codes"] == ["NO_RUNTIME_CURVE"], "ordinate")
    require(artifact["schema_envelope_status"]["runtime_publication_allowed"] is False, "publishable")
    require(artifact["pathway"]["pathway_id"] == "ice_accretion" and artifact["pathway"]["fallback_allowed"] is False, "pathway")
    require(set(artifact["pathway"]["deferred_neighboring_pathways"]) == DEFERRED, "deferred set")
    require({u["id"] for u in artifact["failure_units"]} == UNITS, "failure units")
    require(all(u["ordinate_status"] == "withheld" for u in artifact["failure_units"]), "unit not withheld")
    require(capability == artifact["capability_declaration"], "capability mismatch")
    require(set(capability["metrics_supportable"]) == METRICS and set(capability["metrics_supportable"].values()) == {"withheld"},
            "metrics not withheld")
    require(all("NO_RUNTIME_CURVE" in capability["withheld_reason_by_metric"][m] for m in METRICS), "reason missing")
    require(capability["spread_carried"] is False and capability["emit_modes_populated_by_cell"] == [], "emit")
    index = (ROOT / "docs/contracts/machine_readable_artifact_index.json").read_text()
    require("winter_storm" not in index, "leaked into index")
    require("winter_storm" not in (ROOT / "docs/cells/VERSION_REGISTRY.md").read_text(), "leaked into registry")
    for key in ("source_dossier", "source_workbook", "source_register", "claim_parameter_register", "bounded_evidence_search_log",
                "parameter_tier_table_ref", "value_crosswalk_ref", "legacy_evidence_ingestion_ref", "pressure_test_ref", "known_answer_tests"):
        require((ROOT / artifact[key]).exists(), f"{key} missing")
    require((ROOT / "docs/cells/winter_storm_wind/README.md").exists(), "cell README")
    for doc in DOCS:
        path = P / f"{doc}_winter_storm_wind__{TAG}.md"
        require(path.exists() and path.read_text().strip(), f"missing {path.name}")
    require((P / f"winter_storm_wind_damage_code_metadata_spec__{TAG}.md").exists(), "metadata spec")

    sources = rows(SRC)
    ids = {s["source_id"] for s in sources}
    require(len(ids) == len(sources) and all(s["evidence_tier"] in TIERS for s in sources), "sources")
    legacy = [s for s in sources if s["source_id"] == "WSW-S019"]
    require(legacy and legacy[0]["decision"] == "reject", "legacy rejection not recorded")
    for claim in rows(CLM):
        require(claim["evidence_tier"] in TIERS, "claim tier")
        for sid in claim["source_ids"].split(";"):
            require(sid in ids, f"{claim['claim_id']}: orphan {sid}")
    for param in rows(PAR):
        require(param["tier"] in TIERS, "param tier")
        for sid in param["source_ids"].split(";"):
            require(sid in ids, f"{param['parameter']}: orphan {sid}")

    values = rows(VAL)
    installed = sum(float(v["value"]) for v in values)
    physical = sum(float(v["value"]) for v in values if v["financial_class"] != "excluded_soft_sunk")
    require(math.isclose(installed, 1968, abs_tol=1e-9) and math.isclose(physical, 1623, abs_tol=1e-9), "value reconciliation")
    require(not any(v["role_in_loss"] == "direct_vulnerable_value" for v in values), "direct value enabled in scaffold")

    kats = json.loads(KAT.read_text())
    require(kats["runtime_curve_count"] == 0, "KAT runtime count")
    for test in kats["tests"]:
        got = scaffold_response(test["input"])
        if "expected_error_code" in test:
            require(got.get("error") == test["expected_error_code"], f"{test['test_id']}: {got}")
        else:
            require(got == test["expected_output"], f"{test['test_id']}: numeric output")

    require(load_workbook(WB, read_only=True).sheetnames ==
            ["README", "Failure_Units", "Legacy_Audit", "Value_Crosswalk", "Parameter_Tiers", "Claim_Register", "Source_Register",
             "QA_Checks"], "workbook sheets")
    print(f"winter_storm_wind v0.1 scaffold: PASS ({CHECKS} checks; {len(kats['tests'])} KATs)")
    print("artifact sha256", hashlib.sha256(ART.read_bytes()).hexdigest())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
