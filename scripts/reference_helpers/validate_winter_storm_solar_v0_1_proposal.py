#!/usr/bin/env python3.12
"""Validate the noncanonical winter_storm_solar model v0.1 / docs r1 screening-proxy proposal."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import sys
from pathlib import Path

import jsonschema
from openpyxl import load_workbook
from referencing import Registry, Resource

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from pathway_damage_curve_eval import PathwayEvaluationError  # noqa: E402
from validate_runtime_contracts import validate_v3_curve_record  # noqa: E402
from winter_storm_solar_curve_eval import assemble_loss, evaluate  # noqa: E402

ROOT = HERE.parents[1]
P = ROOT / "docs/cells/winter_storm_solar/proposed"
TAG = "model_v0_1__docs_r1"
ART = P / f"winter_storm_solar__{TAG}__curve_artifact.json"
CAP = P / f"winter_storm_solar__{TAG}__capability.json"
KAT = P / f"known_answer_tests_winter_storm_solar__{TAG}.json"
SRC = P / f"SOURCE_REGISTER_winter_storm_solar__{TAG}.csv"
CLM = P / f"CLAIM_PARAMETER_REGISTER_winter_storm_solar__{TAG}.csv"
PAR = P / f"PARAMETER_TIER_TABLE_winter_storm_solar__{TAG}.csv"
VAL = P / f"VALUE_CROSSWALK_winter_storm_solar__{TAG}.csv"
WB = P / f"damage_curve_records_winter_storm_solar__{TAG}.xlsx"
REQUIRED_DOCS = [
    ROOT / "docs/cells/winter_storm_solar/README.md",
    P / f"README_winter_storm_solar__{TAG}.md",
    P / f"CHANGE_CLASSIFICATION_winter_storm_solar__{TAG}.md",
    P / f"DECISION_LOG_winter_storm_solar__{TAG}.md",
    P / f"BOUNDED_EVIDENCE_SEARCH_LOG_winter_storm_solar__{TAG}.md",
    P / f"PRESSURE_TEST_winter_storm_solar__{TAG}.md",
    P / f"SEVEN_STEP_AUDIT_winter_storm_solar__{TAG}.md",
    P / f"SITE_CONDITION_ADAPTER_winter_storm_solar__{TAG}.md",
    P / f"winter_storm_solar_curve_derivation_dossier__{TAG}.md",
    P / f"winter_storm_solar_damage_code_metadata_spec__{TAG}.md",
    P / f"PROMOTION_GATE_MATRIX_winter_storm_solar__{TAG}.md",
    P / f"VALIDATION_REPORT_winter_storm_solar__{TAG}.md",
]
TIERS = {"T1_claims_or_field_calibrated", "T2_public_lab_standard_or_physics",
         "T3_engineering_proxy_or_adjacent_empirical", "T4_placeholder_or_expert_judgment"}
SHEETS = ["README", "Structure_Curve", "Module_Direct_Curve", "Rack_Tests", "Value_Crosswalk", "Parameter_Tiers",
          "Claim_Register", "Source_Register", "Known_Answers", "QA_Checks"]
CHECKS = 0


def require(condition: bool, message: str) -> None:
    global CHECKS
    CHECKS += 1
    if not condition:
        raise AssertionError(message)


def rows(path: Path) -> list[dict]:
    raw = path.read_bytes()
    require(b"\r\n" not in raw, f"{path.name}: CRLF line endings")
    with path.open(newline="") as handle:
        reader = csv.DictReader(handle)
        out = list(reader)
    require(bool(out), f"{path.name}: empty")
    for line, row in enumerate(out, start=2):
        require(None not in row and all(v is not None for v in row.values()), f"{path.name}:{line}: ragged row")
    return out


def close(a: float, b: float, tol: float) -> bool:
    return math.isclose(a, b, rel_tol=0.0, abs_tol=tol)


def main() -> int:
    artifact = json.loads(ART.read_text())
    capability = json.loads(CAP.read_text())
    schemas = ROOT / "docs/contracts/schemas"
    registry = Registry().with_resources(
        (doc["$id"], Resource.from_contents(doc))
        for doc in (json.loads(path.read_text()) for path in schemas.glob("*.json"))
        if "$id" in doc
    )
    bundle_schema = json.loads((schemas / "curve_artifact_bundle.v3.schema.json").read_text())
    jsonschema.Draft202012Validator(bundle_schema, registry=registry).validate(artifact)
    capability_schema = json.loads((schemas / "capability_declaration.v3.schema.json").read_text())
    jsonschema.Draft202012Validator(capability_schema, registry=registry).validate(capability)
    require(True, "schemas")

    # identity and non-promotion
    require(artifact["cell_id"] == "winter_storm_solar", "cell_id")
    require(artifact["semantic_damage_model_version"] == "model v0.1", "model version")
    require(artifact["documentation_revision"] == "docs r1", "docs revision")
    require(artifact["canonical_runtime_artifact"] is False and capability["canonical_runtime_artifact"] is False, "became canonical")
    require(artifact["package_inclusion_status"] == "not_included" and artifact["package_release"] == "unreleased", "package status")
    require(artifact["promotion_status"] == "proposed", "promotion status")
    require(capability == artifact["capability_declaration"], "embedded and standalone capability differ")
    index = (ROOT / "docs/contracts/machine_readable_artifact_index.json").read_text()
    require("winter_storm" not in index, "proposal leaked into the artifact index")
    require("winter_storm" not in (ROOT / "docs/cells/VERSION_REGISTRY.md").read_text(), "proposal leaked into VERSION_REGISTRY")
    for key in ("source_dossier", "source_workbook", "known_answer_tests", "source_register", "claim_parameter_register", "value_crosswalk"):
        require((ROOT / artifact[key]).exists(), f"{key} path missing")
    for doc in REQUIRED_DOCS:
        require(doc.exists() and doc.read_text().strip(), f"missing doc {doc.relative_to(ROOT)}")

    # curve records and coverage
    pathway = artifact["pathways"][0]
    require([p["pathway_id"] for p in artifact["pathways"]] == ["snow_gravity_load"], "pathway set")
    require([p["pathway_id"] for p in capability["pathway_capabilities"]] == ["snow_gravity_load"], "capability pathways")
    records = pathway["curve_records"]
    require(len(records) == 2, "two runtime records")
    for record in records:
        validate_v3_curve_record(record)
        require(record["pathway_id"] == "snow_gravity_load", "record pathway")
        require([s["scenario_id"] for s in record["parameters"]["capacity_scenarios"]] ==
                ["lower_resistance", "central_screening", "upper_resistance"], "scenario identities")
    medians = [s["state_medians"][0] for s in records[0]["parameters"]["capacity_scenarios"]]
    require(medians == sorted(medians), "structure scenarios ordered weak to strong")
    unit_ids = {f["id"] for f in artifact["failure_units"]}
    coverage = {c["failure_unit_id"]: c["status"] for c in pathway["failure_unit_coverage"]}
    require(set(coverage) == unit_ids, "coverage does not tile failure units")
    curve_units = {r["failure_unit_id"] for r in records}
    require(curve_units == {u for u, s in coverage.items() if s == "conditional_screening_curve"}, "curve/coverage mismatch")
    withheld_cap = {w["failure_unit_id"] for w in capability["pathway_capabilities"][0]["withheld_failure_units"]}
    require(withheld_cap == unit_ids - curve_units, "capability withheld set mismatch")
    require(capability["pathway_capabilities"][0]["failure_unit_scalar_dr"] == "conditional", "DR must stay conditional")

    # registers
    sources = rows(SRC)
    source_ids = {r["source_id"] for r in sources}
    require(len(source_ids) == len(sources), "duplicate source ids")
    require(all(r["evidence_tier"] in TIERS for r in sources), "noncanonical source tier")
    allowed_extra = {"BOUNDED_SEARCH_LOG"}
    claims = rows(CLM)
    require(len({c["claim_id"] for c in claims}) == len(claims), "duplicate claim ids")
    for claim in claims:
        require(claim["evidence_tier"] in TIERS, f"{claim['claim_id']}: tier")
        for sid in claim["source_ids"].split(";"):
            require(sid in source_ids | allowed_extra, f"{claim['claim_id']}: orphan source {sid}")
    params = rows(PAR)
    for param in params:
        require(param["tier"] in TIERS, f"{param['parameter']}: tier")
        for sid in param["source_ids"].split(";"):
            require(sid in source_ids, f"{param['parameter']}: orphan source {sid}")
    for record in records:
        for state in record["parameters"]["damage_states"] + record["parameters"]["capacity_scenarios"]:
            require(state["tier"] in TIERS, "record tier")
            require(set(state["source_ids"]) <= source_ids, "record orphan source")

    # value reconciliation
    values = rows(VAL)
    physical = sum(float(v["value"]) for v in values if v["financial_class"] != "excluded_soft_sunk")
    installed = sum(float(v["value"]) for v in values)
    require(close(physical, 877.7957023626668, 1e-9), f"physical basis {physical}")
    require(close(installed, 1120.0, 1e-9), f"installed basis {installed}")
    linkage = artifact["value_linkage"]
    support = sum(float(v["value"]) for v in values if v["role_in_loss"] == "support_once")
    require(close(support, linkage["support_allocation_rule"]["pool_usd_per_kwdc"], 1e-9), "support pool mismatch")
    direct_hw = sum(float(v["value"]) for v in values if v["financial_class"] == "direct_hardware")
    require(close(direct_hw, linkage["support_allocation_rule"]["direct_hardware_basis_usd_per_kwdc"], 1e-9), "direct hardware basis")
    require(not any(v["role_in_loss"] == "direct_vulnerable_value" and v["failure_unit_id"] not in curve_units for v in values),
            "direct value mapped to a unit without a curve")

    # known answers through the evaluator (independent of the builder's closed form)
    kats = json.loads(KAT.read_text())
    tol = kats["absolute_tolerance"]
    for test in kats["runtime_curve_known_answer_tests"]:
        got = evaluate(artifact, test["input"])
        exp = test["expected"]
        require(got["curve_id"] == exp["curve_id"], f"{test['test_id']}: curve")
        require(close(got["axis"]["ratio"], exp["axis_ratio"], tol), f"{test['test_id']}: ratio")
        if "axis_array_load_kpa" in exp:
            require(close(got["axis"]["array_load_kpa"], exp["axis_array_load_kpa"], tol), f"{test['test_id']}: array load")
        for sid, value in exp["scenario_drs"].items():
            require(close(got["scenario_drs"][sid], value, tol), f"{test['test_id']}: {sid} {got['scenario_drs'][sid]} != {value}")
        require(set(exp["metadata_flags_include"]) <= set(got["flags"]), f"{test['test_id']}: flags {got['flags']}")
    for test in kats["contract_rejection_tests"]:
        try:
            evaluate(artifact, test["input"])
        except PathwayEvaluationError as error:
            require(error.code == test["expected_error_code"], f"{test['test_id']}: {error.code}")
        else:
            require(False, f"{test['test_id']}: accepted")
    for test in kats["loss_assembly_known_answer_tests"]:
        got = assemble_loss(artifact, test["structure_dr_central"], test["module_dr_central"], test["input_values"])
        for key, value in test["expected"].items():
            require(close(got[key], value, kats["monetary_absolute_tolerance"]), f"{test['test_id']}: {key}")
    for test in kats["loss_assembly_rejection_tests"]:
        try:
            assemble_loss(artifact, 0.5, 0.5, test["input_values"])
        except PathwayEvaluationError as error:
            require(error.code == test["expected_error_code"], f"{test['test_id']}: {error.code}")
        else:
            require(False, f"{test['test_id']}: accepted")
    # monotonicity and bounds
    for r in [i * 0.05 for i in range(0, 81)]:
        payload = {"pathway_id": "snow_gravity_load", "failure_unit_id": "WSS1_SUPPORT_STRUCTURE_COLLAPSE",
                   "array_architecture": "fixed_tilt_ground_mount", "event_to_design_snow_demand_ratio": r}
        drs = evaluate(artifact, payload)["scenario_drs"]
        require(drs["lower_resistance"] >= drs["central_screening"] >= drs["upper_resistance"], f"scenario order at r={r}")
        require(all(0 <= v <= 1 for v in drs.values()), "DR bounds")

    # workbook
    wb = load_workbook(WB, read_only=True)
    require(wb.sheetnames == SHEETS, f"workbook sheets {wb.sheetnames}")

    print(f"winter_storm_solar v0.1 proposal: PASS ({CHECKS} checks; "
          f"{len(kats['runtime_curve_known_answer_tests'])} runtime KATs, {len(kats['contract_rejection_tests'])} rejections)")
    print("artifact sha256", hashlib.sha256(ART.read_bytes()).hexdigest())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
