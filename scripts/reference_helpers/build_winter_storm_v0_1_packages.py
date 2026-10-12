#!/usr/bin/env python3.12
"""Build the proposed winter_storm_solar (screening proxy) and winter_storm_wind (fail-closed scaffold)
model v0.1 / docs r1 machine files: artifacts, capabilities, known-answer tests, registers, workbooks.

Known-answer expectations are computed here with closed-form arithmetic that does not import the reference
evaluator, so the validator's evaluator run is an independent check. Deterministic output; rerun after edits.
"""

from __future__ import annotations

import csv
import json
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

import winter_storm_v0_1_registers as R

ROOT = Path(__file__).resolve().parents[2]
SOLAR_DIR = ROOT / "docs/cells/winter_storm_solar/proposed"
WIND_DIR = ROOT / "docs/cells/winter_storm_wind/proposed"
TAG = "model_v0_1__docs_r1"


def rel(path: Path) -> str:
    return str(path.relative_to(ROOT))


def sp(name: str) -> Path:
    return SOLAR_DIR / name.format(cell="winter_storm_solar", tag=TAG)


def wp(name: str) -> Path:
    return WIND_DIR / name.format(cell="winter_storm_wind", tag=TAG)


def write_json(path: Path, obj) -> None:
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=True) + "\n")


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        writer.writerows(rows)


# ================================================================================ SOLAR
S_ART = sp("{cell}__{tag}__curve_artifact.json")
S_CAP = sp("{cell}__{tag}__capability.json")
S_KAT = sp("known_answer_tests_{cell}__{tag}.json")
S_SRC = sp("SOURCE_REGISTER_{cell}__{tag}.csv")
S_CLM = sp("CLAIM_PARAMETER_REGISTER_{cell}__{tag}.csv")
S_PAR = sp("PARAMETER_TIER_TABLE_{cell}__{tag}.csv")
S_VAL = sp("VALUE_CROSSWALK_{cell}__{tag}.csv")
S_WB = sp("damage_curve_records_{cell}__{tag}.xlsx")

G = 9.80665 / 1000.0
GLAZE = math.pi * 0.9 * G
CE, CT, FLAT = 0.9, 1.2, 0.7
STRUCT_BETA, MODULE_BETA = 0.47, 0.25
STRUCT_SCEN = {"lower_resistance": [1.05, 1.26], "central_screening": [1.32, 1.584], "upper_resistance": [1.88, 2.256]}
STRUCT_COST = [0.0, 0.40, 1.00]
MOD_SCEN = {"lower_resistance": [1.8 * 0.85], "central_screening": [2.3 * 0.85], "upper_resistance": [2.8 * 0.85]}
MOD_COST = [0.0, 1.0]
NONSALVAGE = 0.8
SUPPORT_POOL = 43.273972602739725 + 39.447761194029844 + 106.86876712328767
DIRECT_HW = sum(r[3] for r in R.SOLAR_VALUE_ROWS if r[4] == "direct_hardware")
STRUCT_ID = "WSS1_SUPPORT_STRUCTURE_ORDERED_STATES"
MOD_ID = "WSS1_MODULE_FIELD_DIRECT_ORDERED_STATES"


def _phi(z: float) -> float:
    return 0.5 * (1.0 + math.erf(z / math.sqrt(2.0)))


def _dr(x: float, medians: list[float], costs: list[float], beta: float) -> float:
    """Closed form: E[DR] = sum_k c_k (P_k - P_{k+1}); independent of the evaluator."""
    if x <= 0:
        return 0.0
    exceed = [_phi(math.log(x / m) / beta) for m in medians] + [0.0]
    total = costs[0] * (1 - exceed[0])
    for k in range(1, len(costs)):
        total += costs[k] * (exceed[k - 1] - exceed[k])
    return min(1.0, max(0.0, total))


def solar_axes(swe: float, ice: float, design_ground: float | None = None, rated_array: float | None = None, cs: float = 1.0):
    ground = swe * G + ice * GLAZE
    q = cs * FLAT * CE * CT * ground
    ratio = ground / design_ground if design_ground else q / rated_array
    return ground, q, ratio


def struct_drs(r: float) -> dict[str, float]:
    return {s: _dr(r, m, STRUCT_COST, STRUCT_BETA) for s, m in STRUCT_SCEN.items()}


def module_drs(q: float, r: float) -> dict[str, float]:
    s = struct_drs(r)
    return {k: 1 - (1 - _dr(q, MOD_SCEN[k], MOD_COST, MODULE_BETA)) * (1 - NONSALVAGE * s[k]) for k in MOD_SCEN}


def solar_records() -> list[dict]:
    return [
        {
            "curve_id": STRUCT_ID, "pathway_id": "snow_gravity_load", "failure_unit_id": "WSS1_SUPPORT_STRUCTURE_COLLAPSE",
            "curve_form": "ordered_damage_state_lognormal", "x_axis": "event_to_design_snow_demand_ratio",
            "y_axis": "failure_unit_damage_ratio",
            "parameters": {
                "beta_ln": STRUCT_BETA,
                "damage_states": [
                    {"state_id": "STRUCT_DS0_NO_STRUCTURAL_DAMAGE", "cost_ratio": 0.0, "affected_subsystems": [],
                     "description": "no structural repair", "tier": R.T4, "source_ids": ["WSS-S022"]},
                    {"state_id": "STRUCT_DS1_PARTIAL_ARRAY_COLLAPSE", "cost_ratio": 0.40,
                     "affected_subsystems": ["racking or tracker members", "joints"],
                     "description": "member buckling or joint failure collapsing part of the array (~37% field case)",
                     "tier": R.T3, "source_ids": ["WSS-S030", "WSS-S023"]},
                    {"state_id": "STRUCT_DS2_EXTENSIVE_COLLAPSE", "cost_ratio": 1.0,
                     "affected_subsystems": ["racking or tracker support structure"],
                     "description": "extensive to total collapse (~80% and total field cases)",
                     "tier": R.T3, "source_ids": ["WSS-S030", "WSS-S013"]},
                ],
                "capacity_scenarios": [
                    {"scenario_id": "lower_resistance", "state_medians": STRUCT_SCEN["lower_resistance"], "tier": R.T4,
                     "source_ids": ["WSS-S022", "WSS-S023", "WSS-S019"],
                     "interpretation": "test design load read as the old 50-year nominal load (x 0.56); not a percentile"},
                    {"scenario_id": "central_screening", "state_medians": STRUCT_SCEN["central_screening"], "tier": R.T4,
                     "source_ids": ["WSS-S022", "WSS-S023", "WSS-S002"],
                     "interpretation": "test design load read as ASD 0.7 pg (x 0.7); screening midpoint, not fitted to US data"},
                    {"scenario_id": "upper_resistance", "state_medians": STRUCT_SCEN["upper_resistance"], "tier": R.T4,
                     "source_ids": ["WSS-S022", "WSS-S023"],
                     "interpretation": "test design load read as the ASCE 7-22 pg (x 1.0); not a percentile"},
                ],
            },
        },
        {
            "curve_id": MOD_ID, "pathway_id": "snow_gravity_load", "failure_unit_id": "WSS1_MODULE_FIELD_BREAKAGE",
            "curve_form": "ordered_damage_state_lognormal", "x_axis": "array_snow_load_kpa",
            "y_axis": "failure_unit_damage_ratio",
            "parameters": {
                "beta_ln": MODULE_BETA,
                "damage_states": [
                    {"state_id": "MODULE_DS0_NO_BREAKAGE", "cost_ratio": 0.0, "affected_subsystems": [],
                     "description": "no direct module breakage", "tier": R.T4, "source_ids": ["WSS-S011"]},
                    {"state_id": "MODULE_DS1_BREAKAGE_OR_FRAME_FAILURE", "cost_ratio": 1.0,
                     "affected_subsystems": ["modules", "frames", "clamps"],
                     "description": "glass breakage or frame bending/pull-out requiring replacement of the affected modules",
                     "tier": R.T4, "source_ids": ["WSS-S003", "WSS-S030"]},
                ],
                "capacity_scenarios": [
                    {"scenario_id": k, "state_medians": v, "tier": R.T3, "source_ids": ["WSS-S011", "WSS-S010", "WSS-S003"],
                     "interpretation": f"module failure load {round(v[0] / 0.85, 2)} kPa x 0.85 lower-edge knock-down (T4); not a percentile"}
                    for k, v in MOD_SCEN.items()
                ],
            },
        },
    ]


def solar_failure_units() -> list[dict]:
    fu = lambda i, sub, comp, treat, y, den, **kw: {"id": i, "subsystem": sub, "component": comp, "treatment": treat,
                                                     "y_axis": y, "denominator": den, **kw}
    return [
        fu("WSS1_SUPPORT_STRUCTURE_COLLAPSE", "MOUNTING", "TRACKER + RACKING_STRUCTURE excluding foundations", "primary_nonzero",
           "same-unit direct repair or replacement cost divided by support-structure replacement value",
           "explicit site structure value; repository reference 109.98972602739727 2024 USD/kWdc is informational only",
           mapping="subsystem_only"),
        fu("WSS1_MODULE_FIELD_BREAKAGE", "PV_ARRAY", "PV_MODULE (modules, frames, clamps)", "primary_nonzero",
           "same-unit direct replacement cost divided by module-field replacement value",
           "explicit site module value; repository reference 291.21485143992487 2024 USD/kWdc is informational only",
           mapping="exact_component"),
        fu("WSS1_FOUNDATION", "FOUNDATION", "FOUNDATION_BASE", "withheld", "no runtime curve", "not applicable",
           withheld_reason_codes=["NO_SNOW_LOAD_FOUNDATION_RESPONSE_EVIDENCE", "FROST_HEAVE_IS_SEPARATE_CHRONIC_PATHWAY"]),
        fu("WSS1_POWER_CONVERSION_AND_ELECTRICAL", "INVERTER_SYSTEM; ELECTRICAL_COLLECTION; SUBSTATION; GROUNDING_LIGHTNING",
           "INVERTER; COMBINER_BOX; CABLE_AC_DC; TRANSFORMER_MAIN; SWITCHGEAR; GROUNDING_GRID", "withheld", "no runtime curve",
           "not applicable", withheld_reason_codes=["NO_SNOW_LOAD_RESPONSE_EVIDENCE", "ICE_ON_YARD_EQUIPMENT_IS_SEPARATE_PATHWAY"]),
        fu("WSS1_SCADA_COMMUNICATIONS", "SCADA", "MONITORING_SYSTEM", "withheld", "no runtime curve", "not applicable",
           withheld_reason_codes=["NO_SNOW_LOAD_RESPONSE_EVIDENCE"]),
        fu("WSS1_CIVIL_INFRA", "CIVIL_INFRA", "SITE_PREP; buildings; fencing; roads", "withheld", "no runtime curve",
           "not applicable", withheld_reason_codes=["MIXED_ASSET_BUCKET_REQUIRES_SPLIT", "BUILDING_ROOF_SNOW_IS_SEPARATE_ASSET_CLASS"]),
        fu("WSS1_REPLACEMENT_SUPPORT", "replacement support", "SBOS/EBOS field labor, site management, equipment rental",
           "exposure_modifier", "no intrinsic damage ratio",
           f"{SUPPORT_POOL} 2024 USD/kWdc reference pool; allocate once after direct damage",
           withheld_reason_codes=["ALLOCATION_ONLY_NO_INTRINSIC_DR", "ALLOCATE_ONCE_AFTER_DIRECT_DAMAGE"]),
    ]


def solar_capability() -> dict:
    return {
        "schema_version": "capability_declaration.v3",
        "cell_id": "winter_storm_solar",
        "canonical_runtime_artifact": False,
        "pathway_capabilities": [{
            "pathway_id": "snow_gravity_load",
            "failure_unit_scalar_dr": "conditional",
            "scenario_loss_given_value_basis": "conditional",
            "curve_intrinsic_spread": "nonprobabilistic_epistemic_envelope",
            "populated_emit_modes": ["scalar_mean_plus_bounds", "state_ensemble"],
            "conditions": [
                "pathway_id snow_gravity_load and a supported ground-mount architecture supplied",
                "snowpack water equivalent (not depth) supplied with an explicit ice_component_basis, or a direct demand ratio",
                "racking design basis declared: site_code_assumed with the cell design ground snow load, or project_rated with the rated array snow load",
                "module and structure values and exposure supplied for monetary loss",
                "module-structure cascade and support allocation applied once",
            ],
            "limitation_flags": [
                "SCREENING_ENGINEERING_PROXY",
                "NONPROBABILISTIC_EPISTEMIC_ENVELOPE",
                "CAPACITY_FROM_SEVEN_JAPANESE_FIXED_RACK_TESTS",
                "JAPANESE_TO_ASCE_DESIGN_LOAD_MAPPING_UNRESOLVED",
                "TRACKER_TRANSFER_FROM_FIXED_RACK_TESTS",
                "SNOW_STOW_AND_SNOW_REMOVAL_NO_CREDIT",
                "MODULE_CAPACITY_T3_LOWER_EDGE_KNOCKDOWN_T4",
                "MODULE_SALVAGE_T4_ASSUMPTION",
                "NO_US_FIELD_OR_CLAIMS_CALIBRATION",
                "FULL_PLANT_PHYSICAL_LOSS_INCOMPLETE",
            ],
            "withheld_failure_units": [
                {"failure_unit_id": f["id"], "reason_codes": f["withheld_reason_codes"]}
                for f in solar_failure_units() if f["treatment"] != "primary_nonzero"
            ],
        }],
        "consumer_annual_metrics": {
            "computation_owner": "downstream_consumer",
            "status_before_promotion": "withheld_noncanonical_proposal",
            "status_after_promotion": "consumer_computable_from_validated_frequency_intensity_coupling_value_and_cap_model",
            "prerequisites": [
                "canonical promotion and exact model/docs/schema/SHA pin",
                "winter event catalog with peak snowpack SWE, ice thickness and parent event_family_id",
                "per-cell or per-asset design snow load",
                "architecture-aware exposure and value basis",
                "consumer cap-binding validation",
            ],
            "limitation_flags": [
                "FREQUENCY_DRIVEN_TAIL_IS_CONSUMER_OWNED",
                "ENGINEERING_SCENARIOS_ARE_NOT_A_PROBABILITY_DISTRIBUTION",
                "FULL_PLANT_LOSS_REQUIRES_SEPARATELY_QUALIFIED_FAILURE_UNITS",
            ],
        },
        "cap_binding": {
            "policy": "consumer_enforced_fail_closed",
            "enforcement_owner": "downstream_consumer",
            "checks_required": [
                "exact model/docs/schema/SHA pin",
                "pathway and architecture match",
                "SWE and ice basis provenance",
                "declared racking design basis",
                "explicit values and exposure",
                "module-structure cascade applied once",
                "support costs not embedded and re-added",
                "withheld units remain null",
            ],
            "action_if_fail": "withhold the affected DR, loss, or annual metric; do not substitute another winter or wind pathway",
        },
        "promotion_gate": {
            "status": "blocked",
            "required_before_canonical_use": [
                "independent engineering review of axis, bridge, scenarios and cascade",
                "ASCE 7-22 Ch.7 and the ASCE 2025 ground-mount snow guidance read and reconciled",
                "tracker-specific snow capacity evidence or an explicit tracker transfer review",
                "US field, forensic or claims evidence sufficient to review the T4 mapping",
                "Hazard dual-read and rejection-path tests",
                "explicit maintainer promotion decision and atomic registry/index/changelog update",
            ],
        },
    }


def solar_artifact() -> dict:
    cap = solar_capability()
    return {
        "schema_version": "damage_curve_record_bundle.v3",
        "schema_status": "proposed_draft",
        "cell_id": "winter_storm_solar",
        "damage_code_id": "WINTER_STORM_SOLAR_SNOW_LOAD_V0_1_PROPOSED",
        "semantic_damage_model_version": "model v0.1",
        "documentation_revision": "docs r1",
        "lifecycle_state": "candidate",
        "promotion_status": "proposed",
        "review_status": "pressure_tested_screening_proxy",
        "model_grade": "screening_engineering_proxy",
        "package_release": "unreleased",
        "package_baseline": "library v2.5",
        "package_inclusion_status": "not_included",
        "canonical_runtime_artifact": False,
        "source_dossier": rel(sp("{cell}_curve_derivation_dossier__{tag}.md")),
        "source_workbook": rel(S_WB),
        "known_answer_tests": rel(S_KAT),
        "source_register": rel(S_SRC),
        "claim_parameter_register": rel(S_CLM),
        "value_crosswalk": rel(S_VAL),
        "failure_units": solar_failure_units(),
        "pathways": [{
            "pathway_id": "snow_gravity_load",
            "hazard_scope": {
                "included": ["accumulated snowpack weight on the array, including snow already on the ground",
                             "rain-on-snow water retained in the pack (carried by SWE)",
                             "glaze ice weight on module surfaces, as water-equivalent load"],
                "excluded": ["snowmelt or rain-on-snow flooding", "blizzard or synoptic wind loading", "ice plus concurrent wind",
                             "ice or snow shedding impact", "snow drift and sliding impact on neighbouring equipment",
                             "frost heave and chronic foundation movement", "extreme cold, freeze-off, derating",
                             "hail", "business interruption or downtime"],
                "event_partition_rule": "one parent winter event_family_id; evaluate the peak snowpack reached during the event once; do not add consecutive storms' new snow separately",
            },
            "hazard_axis": {
                "id": "winter_snow_gravity_demand",
                "preferred_input_field": "event_to_design_snow_demand_ratio",
                "permitted_proxy_fields": ["snow_water_equivalent_peak_mm", "ice_radial_thickness_mm",
                                           "design_ground_snow_load_kpa", "racking_rated_array_snow_load_kpa"],
                "unit": "dimensionless",
                "valid_range": [0.0, 4.0],
                "extrapolation_policy": "structure ratio: flag above 3.3, withhold above 4.0; module array load: flag above 4.5 kPa, withhold above 6.0 kPa",
                "axis_contracts": {
                    "structure": {"axis_field": "event_to_design_snow_demand_ratio", "unit": "dimensionless",
                                  "valid_range": [0.0, 4.0], "flag_above": 3.3,
                                  "definition": "event ground frozen load / design ground snow load (site_code_assumed) or event array load / rated array snow load (project_rated)"},
                    "module": {"axis_field": "array_snow_load_kpa", "unit": "kPa", "valid_range": [0.0, 6.0], "flag_above": 4.5,
                               "definition": "Cs x 0.7 x Ce x Ct x event ground frozen load"},
                },
                "load_bridge": {
                    "swe_to_kpa_per_mm": G,
                    "glaze_kpa_per_mm_radial": GLAZE,
                    "flat_roof_factor": FLAT, "exposure_factor_ce": CE, "thermal_factor_ct": CT,
                    "default_slope_factor_cs": 1.0,
                    "source_ids": ["WSS-S031", "WSS-S002", "WSS-S025"],
                    "snow_depth_policy": "reject snow depth without SWE (SNOW_DEPTH_WITHOUT_SWE_UNSUPPORTED)",
                },
            },
            "selector_logic": [
                {"field": "array_architecture", "required": True, "allowed": ["fixed_tilt_ground_mount", "single_axis_tracker"],
                 "no_default": True, "numeric_effect": "none in v0.1; single_axis_tracker adds TRACKER_TRANSFER_FROM_FIXED_RACK_TESTS"},
                {"field": "racking_design_basis", "required": True, "allowed": ["site_code_assumed", "project_rated"],
                 "no_default": True, "numeric_effect": "selects the structure-ratio denominator"},
                {"field": "design_ground_snow_load_kpa", "required_when": "racking_design_basis=site_code_assumed",
                 "numeric_effect": "structure ratio denominator (ASCE 7-22 strength-level pg)"},
                {"field": "racking_rated_array_snow_load_kpa", "required_when": "racking_design_basis=project_rated",
                 "numeric_effect": "structure ratio denominator on the array-load basis"},
                {"field": "design_risk_category", "required": "capture_or_unknown", "allowed": ["I", "II", "unknown"],
                 "numeric_effect": "none; records which ASCE 7-22 map supplied the design load"},
            ],
            "conditioner_logic": [
                {"field": "ice_component_basis", "required": True, "allowed": ["supplied", "not_in_hazard_product"],
                 "numeric_effect": "not_in_hazard_product evaluates snow only and flags ICE_COMPONENT_NOT_SUPPLIED_LOWER_BOUND"},
                {"field": "slope_factor_cs", "required": False, "default": 1.0,
                 "numeric_effect": "module axis only; a value below 1 is flagged SLOPE_FACTOR_CREDIT_SUPPLIED_UNCALIBRATED"},
                {"field": "snow_stow_state", "required": "capture_or_unknown",
                 "allowed": ["confirmed_max_tilt", "not_confirmed", "unknown", "not_applicable_fixed_tilt"],
                 "numeric_effect": "none; no stow credit"},
                {"field": "leading_edge_clearance_m", "required": "capture_or_unknown", "numeric_effect": "none; capture only"},
                {"field": "snow_removal_performed", "required": "capture_or_unknown", "numeric_effect": "none; no removal credit"},
            ],
            "exposure_contract": {
                "required_for_loss": ["event_id", "event_family_id", "array_zone_id_or_group", "exposed_fraction",
                                      "explicit_failure_unit_values"],
                "no_default_exposed_fraction": True,
                "rules": ["evaluate each array zone once per event", "do not multiply by installed TIV",
                          "do not reuse array exposure for foundations, electrical, SCADA or civil units",
                          "reconcile shared module disposition with hail and wind loss atoms before summing"],
            },
            "failure_unit_coverage": [
                {"failure_unit_id": "WSS1_SUPPORT_STRUCTURE_COLLAPSE", "status": "conditional_screening_curve"},
                {"failure_unit_id": "WSS1_MODULE_FIELD_BREAKAGE", "status": "conditional_screening_curve"},
                {"failure_unit_id": "WSS1_FOUNDATION", "status": "withheld"},
                {"failure_unit_id": "WSS1_POWER_CONVERSION_AND_ELECTRICAL", "status": "withheld"},
                {"failure_unit_id": "WSS1_SCADA_COMMUNICATIONS", "status": "withheld"},
                {"failure_unit_id": "WSS1_CIVIL_INFRA", "status": "withheld"},
                {"failure_unit_id": "WSS1_REPLACEMENT_SUPPORT", "status": "allocation_only"},
            ],
            "curve_records": solar_records(),
        }],
        "deferred_pathways": [
            {"pathway_id": "ice_with_concurrent_wind", "reason": "wind on iced or snow-loaded arrays is a combined case (WSS-S027); no evidence chain"},
            {"pathway_id": "ice_shedding_impact", "reason": "falling or sliding ice/snow impact on neighbouring equipment; no evidence chain"},
            {"pathway_id": "snow_drift_and_sliding_impact", "reason": "drift and sliding loads on equipment below arrays; no evidence chain"},
        ],
        "value_linkage": {
            "implicit_default_profile": None,
            "primary_reference_basis": {
                "basis_id": "NLR_Q1_2025_UPV_PV_ONLY_2024_USD",
                "source": "docs/method/value_basis/solar_wind_value_breakdown.xlsx",
                "installed_capex_usd_per_kwdc": 1120.0,
                "physical_replaceable_usd_per_kwdc": 877.7957023626668,
                "direct_hardware_usd_per_kwdc": DIRECT_HW,
                "replacement_support_usd_per_kwdc": SUPPORT_POOL,
                "module_reference_usd_per_kwdc": 291.21485143992487,
                "structure_reference_usd_per_kwdc": 109.98972602739727,
                "currency_basis": "2024 USD",
            },
            "support_allocation_rule": {
                "pool_usd_per_kwdc": SUPPORT_POOL,
                "direct_hardware_basis_usd_per_kwdc": DIRECT_HW,
                "rule": "support_once_loss = pool x (direct snow loss / direct hardware basis); applied once after direct damage",
                "tier": R.T4,
            },
            "loss_requirements": ["explicit module and structure values with currency and price year",
                                  "explicit exposed fraction per array zone", "no installed-TIV ratio"],
        },
        "parameter_tier_table": [dict(zip(R.PARAMETER_FIELDS, p)) for p in R.SOLAR_PARAMETERS],
        "evaluation_contract": {
            "pathway_id_required": True,
            "architecture_required": True,
            "no_cross_pathway_fallback": True,
            "central_scenario_id": "central_screening",
            "scenario_semantics": "unweighted epistemic resistance scenarios, not percentiles or a probability distribution",
            "module_structure_cascade": {
                "rule": "module DR = 1 - (1 - direct module DR) x (1 - nonsalvage x structure DR), per scenario",
                "central_nonsalvage_fraction": NONSALVAGE,
                "bounds": {"full_salvage": 0.0, "no_salvage": 1.0},
                "tier": R.T4,
            },
            "unsupported_architectures": ["roof_mount", "carport_or_canopy", "floating_pv", "dual_axis_tracker",
                                          "agrivoltaic_elevated_or_vertical", "concentrating_solar"],
            "failure_codes": ["PATHWAY_ID_REQUIRED", "PATHWAY_ID_UNKNOWN", "PATHWAY_DEFERRED_NOT_IMPLEMENTED",
                              "ARRAY_ARCHITECTURE_REQUIRED", "ARRAY_ARCHITECTURE_UNSUPPORTED", "FAILURE_UNIT_ID_UNKNOWN",
                              "NO_RUNTIME_CURVE_FOR_PATHWAY_UNIT", "AXIS_PAYLOAD_REQUIRED", "AXIS_PAYLOAD_INVALID",
                              "SNOW_DEPTH_WITHOUT_SWE_UNSUPPORTED", "ICE_COMPONENT_BASIS_REQUIRED",
                              "RACKING_DESIGN_BASIS_REQUIRED", "DESIGN_SNOW_LOAD_REQUIRED", "AXIS_OUTSIDE_VALID_RANGE",
                              "CONDITIONER_VALUE_UNSUPPORTED", "EXPLICIT_VALUE_AND_EXPOSURE_REQUIRED"],
        },
        "emit_contract": {
            "schema_version": "damage_emit.v2",
            "populated_emit_modes_for_this_cell": ["scalar_mean_plus_bounds", "state_ensemble"],
            "result_grain": "pathway_id x array_architecture x failure_unit_id x array_zone",
            "required_metadata": ["pathway_id", "array_architecture", "racking_design_basis", "ice_component_basis",
                                  "axis values and bridge", "scenario identities", "failure-unit denominator", "limitation flags"],
            "prohibited_outputs": ["full-plant physical DR", "installed-TIV DR", "frequency", "EAL", "PML", "VaR", "TVaR",
                                   "business interruption", "downtime"],
        },
        "capability_declaration": cap,
        "derivation_rationale": {
            "summary": "Seven full-scale rack tests anchor the support-structure capacity spread relative to manufacturer design load; the mapping onto the ASCE 7-22 axis is unresolved and carried as three unweighted scenarios. Module breakage sits on absolute array load from modern test-to-failure data. Damage-state costs come from Japanese regulator cases. All shape choices beyond the measured ratios are T4 screening assumptions.",
            "rejected_models": ["raw kPa axis without design normalisation", "snow-depth axis without SWE",
                                "separate design-deficiency multiplier on top of the measured spread",
                                "stow or snow-removal credit", "single logistic whole-plant curve"],
        },
        "legacy_comparison": {"status": "no legacy winter x solar material exists in the legacy library or this repository"},
    }


def solar_kats() -> dict:
    runtime = []

    def add(test_id, unit, payload, r, q, ground=None, flags=()):
        exp = {"status": "conditional", "axis_ratio": r}
        if unit == "WSS1_SUPPORT_STRUCTURE_COLLAPSE":
            drs = struct_drs(r)
            exp["curve_id"] = STRUCT_ID
        else:
            drs = module_drs(q, r)
            exp["curve_id"] = MOD_ID
            exp["axis_array_load_kpa"] = q
        exp["scenario_drs"] = drs
        exp["scalar_central_dr"] = drs["central_screening"]
        exp["metadata_flags_include"] = list(flags)
        runtime.append({"test_id": test_id, "pathway_id": "snow_gravity_load", "failure_unit_id": unit,
                        "input": {"pathway_id": "snow_gravity_load", "failure_unit_id": unit, **payload}, "expected": exp,
                        "source": "closed-form equation fixture independent of the reference evaluator"})

    add("structure_direct_ratio_1", "WSS1_SUPPORT_STRUCTURE_COLLAPSE",
        {"array_architecture": "fixed_tilt_ground_mount", "event_to_design_snow_demand_ratio": 1.0}, 1.0, None)
    g, q, r = solar_axes(150, 0, design_ground=0.96)
    add("structure_swe150_no_ice_product_tracker", "WSS1_SUPPORT_STRUCTURE_COLLAPSE",
        {"array_architecture": "single_axis_tracker", "snow_water_equivalent_peak_mm": 150, "ice_component_basis": "not_in_hazard_product",
         "racking_design_basis": "site_code_assumed", "design_ground_snow_load_kpa": 0.96}, r, q,
        flags=("ICE_COMPONENT_NOT_SUPPLIED_LOWER_BOUND", "DESIGN_BASIS_ASSUMED_SITE_CODE", "TRACKER_TRANSFER_FROM_FIXED_RACK_TESTS"))
    g, q, r = solar_axes(120, 12, design_ground=0.96)
    add("structure_swe120_ice12_fixed", "WSS1_SUPPORT_STRUCTURE_COLLAPSE",
        {"array_architecture": "fixed_tilt_ground_mount", "snow_water_equivalent_peak_mm": 120, "ice_component_basis": "supplied",
         "ice_radial_thickness_mm": 12, "racking_design_basis": "site_code_assumed", "design_ground_snow_load_kpa": 0.96}, r, q,
        flags=("DESIGN_BASIS_ASSUMED_SITE_CODE",))
    g, q, r = solar_axes(150, 0, design_ground=0.96)
    add("module_swe150_cascade_fixed", "WSS1_MODULE_FIELD_BREAKAGE",
        {"array_architecture": "fixed_tilt_ground_mount", "snow_water_equivalent_peak_mm": 150, "ice_component_basis": "supplied",
         "ice_radial_thickness_mm": 0, "racking_design_basis": "site_code_assumed", "design_ground_snow_load_kpa": 0.96}, r, q,
        flags=("MODULE_SALVAGE_T4_ASSUMPTION",))
    add("module_direct_ratio_and_load", "WSS1_MODULE_FIELD_BREAKAGE",
        {"array_architecture": "fixed_tilt_ground_mount", "event_to_design_snow_demand_ratio": 0.8, "array_snow_load_kpa": 2.0},
        0.8, 2.0, flags=("MODULE_SALVAGE_T4_ASSUMPTION",))
    g, q, r = solar_axes(60, 0, rated_array=0.2394)
    add("structure_tracker_project_rated_5psf", "WSS1_SUPPORT_STRUCTURE_COLLAPSE",
        {"array_architecture": "single_axis_tracker", "snow_water_equivalent_peak_mm": 60, "ice_component_basis": "supplied",
         "ice_radial_thickness_mm": 0, "racking_design_basis": "project_rated", "racking_rated_array_snow_load_kpa": 0.2394}, r, q,
        flags=("TRACKER_TRANSFER_FROM_FIXED_RACK_TESTS",))
    add("structure_zero_snow", "WSS1_SUPPORT_STRUCTURE_COLLAPSE",
        {"array_architecture": "fixed_tilt_ground_mount", "snow_water_equivalent_peak_mm": 0, "ice_component_basis": "supplied",
         "ice_radial_thickness_mm": 0, "racking_design_basis": "site_code_assumed", "design_ground_snow_load_kpa": 0.96}, 0.0, 0.0)
    add("structure_ratio_above_tested_flag", "WSS1_SUPPORT_STRUCTURE_COLLAPSE",
        {"array_architecture": "fixed_tilt_ground_mount", "event_to_design_snow_demand_ratio": 3.5}, 3.5, None,
        flags=("EVENT_TO_DESIGN_SNOW_DEMAND_RATIO_ABOVE_TESTED_RANGE",))

    g, q, r = solar_axes(150, 0, design_ground=0.96)
    s_dr, m_dr = struct_drs(r)["central_screening"], module_drs(q, r)["central_screening"]
    cap_kw, frac = 120000.0, 1.0
    module_loss, struct_loss = m_dr * 291.21485143992487 * cap_kw * frac, s_dr * 109.98972602739727 * cap_kw * frac
    direct = module_loss + struct_loss
    support = SUPPORT_POOL * direct / DIRECT_HW
    loss = [{
        "test_id": "loss_swe150_central_once",
        "source_runtime_test_id": "module_swe150_cascade_fixed",
        "input_values": {"currency": "USD", "price_year": 2024, "value_unit": "currency_per_kwdc",
                         "module_value": 291.21485143992487, "structure_value": 109.98972602739727,
                         "capacity_kwdc": cap_kw, "exposed_fraction": frac, "support_pool_value": SUPPORT_POOL},
        "structure_dr_central": s_dr, "module_dr_central": m_dr,
        "expected": {"module_loss": module_loss, "structure_loss": struct_loss, "direct_array_loss": direct,
                     "support_once_loss": support, "total_loss": direct + support},
        "notes": "Reference values, central scenario, whole canonical array exposed; support pool allocated once.",
    }]
    base = {"pathway_id": "snow_gravity_load", "failure_unit_id": "WSS1_SUPPORT_STRUCTURE_COLLAPSE",
            "array_architecture": "fixed_tilt_ground_mount", "snow_water_equivalent_peak_mm": 100, "ice_component_basis": "supplied",
            "ice_radial_thickness_mm": 0, "racking_design_basis": "site_code_assumed", "design_ground_snow_load_kpa": 1.0}

    def rej(test_id, change, drop, code, note):
        payload = {k: v for k, v in base.items() if k not in drop}
        payload.update(change)
        return {"test_id": test_id, "input": payload, "expected_error_code": code, "notes": note}

    rejections = [
        rej("reject_missing_pathway", {}, ["pathway_id"], "PATHWAY_ID_REQUIRED", "Pathway has no default."),
        rej("reject_deferred_ice_wind", {"pathway_id": "ice_with_concurrent_wind"}, [], "PATHWAY_DEFERRED_NOT_IMPLEMENTED", "Deferred pathway never falls back."),
        rej("reject_unknown_pathway", {"pathway_id": "hail_impact"}, [], "PATHWAY_ID_UNKNOWN", "Neighbouring pathway rejected."),
        rej("reject_missing_architecture", {}, ["array_architecture"], "ARRAY_ARCHITECTURE_REQUIRED", "Architecture has no default."),
        rej("reject_roof_mount", {"array_architecture": "roof_mount"}, [], "ARRAY_ARCHITECTURE_UNSUPPORTED", "Ground mount only."),
        rej("reject_withheld_foundation", {"failure_unit_id": "WSS1_FOUNDATION"}, [], "NO_RUNTIME_CURVE_FOR_PATHWAY_UNIT", "Withheld is not zero."),
        rej("reject_unknown_unit", {"failure_unit_id": "PV_INVERTER"}, [], "FAILURE_UNIT_ID_UNKNOWN", "No alias mapping."),
        rej("reject_snow_depth_only", {"snow_depth_mm": 400}, ["snow_water_equivalent_peak_mm"], "SNOW_DEPTH_WITHOUT_SWE_UNSUPPORTED", "Density unknown."),
        rej("reject_missing_ice_basis", {}, ["ice_component_basis"], "ICE_COMPONENT_BASIS_REQUIRED", "Ice scope must be explicit."),
        rej("reject_missing_design_basis", {}, ["racking_design_basis"], "RACKING_DESIGN_BASIS_REQUIRED", "Design basis has no default."),
        rej("reject_site_code_without_load", {}, ["design_ground_snow_load_kpa"], "DESIGN_SNOW_LOAD_REQUIRED", "Denominator required."),
        rej("reject_negative_swe", {"snow_water_equivalent_peak_mm": -5}, [], "AXIS_OUTSIDE_VALID_RANGE", "Non-negative demand."),
        rej("reject_ratio_above_range", {"event_to_design_snow_demand_ratio": 4.5}, [], "AXIS_OUTSIDE_VALID_RANGE", "Withhold above 4.0."),
        rej("reject_slope_factor_invalid", {"slope_factor_cs": 1.5}, [], "CONDITIONER_VALUE_UNSUPPORTED", "Cs in (0, 1]."),
    ]
    loss_rejections = [{"test_id": "loss_reject_missing_currency", "input_values": {k: v for k, v in loss[0]["input_values"].items() if k != "currency"},
                        "expected_error_code": "EXPLICIT_VALUE_AND_EXPOSURE_REQUIRED", "notes": "Currency is part of the value basis."},
                       {"test_id": "loss_reject_missing_exposure", "input_values": {k: v for k, v in loss[0]["input_values"].items() if k != "exposed_fraction"},
                        "expected_error_code": "EXPLICIT_VALUE_AND_EXPOSURE_REQUIRED", "notes": "No default exposed fraction."}]
    return {"schema_version": "known_answer_tests.v3", "cell_id": "winter_storm_solar", "semantic_damage_model_version": "model v0.1",
            "documentation_revision": "docs r1", "artifact_schema_version": "damage_curve_record_bundle.v3",
            "absolute_tolerance": 1e-12, "monetary_absolute_tolerance": 1e-6,
            "runtime_curve_known_answer_tests": runtime, "loss_assembly_known_answer_tests": loss,
            "loss_assembly_rejection_tests": loss_rejections, "contract_rejection_tests": rejections}


# ================================================================================ WIND
W_ART = wp("{cell}__{tag}__curve_artifact.json")
W_CAP = wp("{cell}__{tag}__capability.json")
W_KAT = wp("known_answer_tests_{cell}__{tag}.json")
W_SRC = wp("SOURCE_REGISTER_{cell}__{tag}.csv")
W_CLM = wp("CLAIM_PARAMETER_REGISTER_{cell}__{tag}.csv")
W_PAR = wp("PARAMETER_TIER_TABLE_{cell}__{tag}.csv")
W_VAL = wp("VALUE_CROSSWALK_{cell}__{tag}.csv")
W_WB = wp("damage_curve_records_{cell}__{tag}.xlsx")
METRICS = ["failure_unit_scalar_dr", "scenario_loss_given_value_basis", "scalar_eal", "pml", "var", "tvar"]

WIND_UNITS = [
    ("WT_BLADE_ASSEMBLY", "ROTOR_ASSEMBLY; PITCH_SYSTEM", "BLADE; PITCH_DRIVE", "candidate_primary_no_curve",
     "ice mass, imbalance and extreme iced-blade loads", "Wind_Map rows 2-3"),
    ("WT_NACELLE_AND_HUB", "NACELLE; ROTOR_ASSEMBLY; POWER_ELECTRONICS", "HUB; NACELLE_STRUCTURE; GEARBOX; GENERATOR_INTERNAL; YAW_SYSTEM; POWER_CONVERTER",
     "withheld_candidate", "ice load transfer; cold effects belong to a separate pathway", "Wind_Map rows 4-8"),
    ("WT_TOWER_AND_EXTERNAL_FIXTURES", "TOWER", "TOWER_SECTION; ladders; cable trays", "withheld_candidate",
     "iced projected area under concurrent wind (deferred pathway)", "Wind_Map row 9"),
    ("WT_PAD_ELECTRICAL", "ELECTRICAL_COLLECTION", "pad transformers and switchgear", "withheld_candidate", "ice on enclosures and bushings", "Wind_Map row 12 (split required)"),
    ("WT_COLLECTION_NETWORK", "ELECTRICAL_COLLECTION", "CABLE_AC (overhead or buried)", "withheld_candidate",
     "ice and ice + wind on overhead spans; buried cables not exposed", "Wind_Map row 12 (split required)"),
    ("WT_GSU_SUBSTATION", "SUBSTATION", "TRANSFORMER_MAIN; SWITCHGEAR; bus", "withheld_candidate", "ice bridging and bus loads", "Wind_Map row 12 (split required)"),
    ("WT_CONTROL_AND_MET_STATION", "SCADA", "met masts; control building", "withheld_candidate", "ice + wind on guyed lattice masts", "not in value workbook"),
    ("WT_FOUNDATION", "FOUNDATION", "FOUNDATION_BASE", "geometry_screened_no_numeric_zero", "no direct ice-load pathway; frost heave separate", "Wind_Map row 10"),
    ("WT_CIVIL_INFRA", "CIVIL_INFRA", "SITE_ACCESS_STAGING_FACILITIES", "split_required", "building roof snow is a separate asset class", "Wind_Map row 11"),
    ("SUPPORT_FIELDWORK", "REPLACEMENT_FIELDWORK", "ASSEMBLY_AND_INSTALLATION", "allocation_only", "allocate once", "Wind_Map row 13"),
    ("SUPPORT_TRANSPORT_LOGISTICS", "REPLACEMENT_LOGISTICS", "WIND_TURBINE_TRANSPORT", "allocation_only", "allocate once", "Wind_Map row 14"),
]


def wind_capability() -> dict:
    reasons = {m: ["NO_RUNTIME_CURVE"] for m in METRICS}
    reasons["scenario_loss_given_value_basis"] += ["MISSING_VALUE_BASIS", "MISSING_EXPOSURE_OR_COUPLING"]
    for m in ("scalar_eal",):
        reasons[m] += ["MISSING_VALUE_BASIS", "MISSING_EXPOSURE_OR_COUPLING", "MISSING_HAZARD_FREQUENCY_OR_INTENSITY_DISTRIBUTION"]
    for m in ("pml", "var", "tvar"):
        reasons[m] += ["MISSING_HAZARD_FREQUENCY_OR_INTENSITY_DISTRIBUTION"]
    return {
        "schema_version": "capability_declaration.v1", "cell_id": "winter_storm_wind", "spread_carried": False,
        "emit_modes_populated_by_cell": [], "metrics_supportable": {m: "withheld" for m in METRICS},
        "withheld_reason_by_metric": reasons,
        "withheld_reason_codes": ["NO_RUNTIME_CURVE", "MISSING_VALUE_BASIS", "MISSING_EXPOSURE_OR_COUPLING",
                                  "MISSING_HAZARD_FREQUENCY_OR_INTENSITY_DISTRIBUTION"],
        "cap_binding": {"policy": "fail_closed", "preflight_status": "not_executed_no_distribution", "required_before_scalar_eal": True,
                        "tolerance_pct": None, "tolerance_status": "not_applicable_no_distribution",
                        "action_if_fail": "withhold_scalar_eal_and_require_tail_capable_emit"},
    }


def wind_artifact() -> dict:
    return {
        "schema_version": "damage_curve_record_bundle.v1",
        "cell_id": "winter_storm_wind",
        "damage_code_id": "WINTER_STORM_WIND_PROPOSED_V0_1",
        "semantic_damage_model_version": "model v0.1",
        "documentation_revision": "docs r1",
        "package_release": "unreleased",
        "package_baseline": "library v2.5",
        "lifecycle_state": "scaffold",
        "promotion_status": "proposed",
        "review_status": "pressure_tested",
        "documentation_status": "working_revision",
        "package_inclusion_status": "not_included",
        "canonical_runtime_artifact": False,
        "schema_envelope_status": {"selected_schema": "damage_curve_record_bundle.v1", "use": "noncanonical_zero_curve_scaffold_only",
                                   "reason": "Repository-current v2 and v3 schemas require output-bearing records; this zero-curve scaffold is not publishable.",
                                   "runtime_publication_allowed": False,
                                   "promotion_rule": "A runtime bundle must use the repository-current schema and contain reviewed output-bearing records."},
        "ordinate_status": "withheld",
        "withheld_reason_codes": ["NO_RUNTIME_CURVE"],
        "source_dossier": rel(wp("{cell}_curve_derivation_dossier__{tag}.md")),
        "source_workbook": rel(W_WB),
        "source_register": rel(W_SRC),
        "claim_parameter_register": rel(W_CLM),
        "bounded_evidence_search_log": rel(wp("BOUNDED_EVIDENCE_SEARCH_LOG_{cell}__{tag}.md")),
        "parameter_tier_table_ref": rel(W_PAR),
        "value_crosswalk_ref": rel(W_VAL),
        "legacy_evidence_ingestion_ref": rel(wp("LEGACY_EVIDENCE_INGESTION_{cell}__{tag}.md")),
        "pressure_test_ref": rel(wp("PRESSURE_TEST_{cell}__{tag}.md")),
        "known_answer_tests": rel(W_KAT),
        "pathway": {
            "pathway_id": "ice_accretion", "physical_mechanism": "atmospheric glaze or rime accretion load on turbine and BOP members",
            "exact_match_required": True, "fallback_allowed": False, "event_family_id_required": True,
            "deferred_neighboring_pathways": ["ice_with_concurrent_wind", "ice_shedding_impact", "snow_gravity_load_on_buildings",
                                              "extreme_cold_physical_effects"],
            "compound_event_rule": "Preserve one parent winter event_family_id; reconcile any shared physical state and value once downstream.",
        },
        "asset_archetype": {"asset": "land_based_multi_megawatt_horizontal_axis_wind_facility",
                            "reference_configuration": "generic_modern_onshore_reference_value_archetype",
                            "target_applicability_status": "not_frozen",
                            "scope_included": ["direct_occurrence_ice_accretion_physical_destruction"],
                            "scope_excluded": ["icing_downtime_and_curtailment", "freeze_off_and_cold_shutdown", "ice_throw_or_shedding",
                                               "ice_plus_concurrent_wind", "snow_on_buildings", "frost_heave", "business_interruption",
                                               "frequency_and_annual_or_tail_metrics"]},
        "value_basis": {"status": "reference_archetype_only_not_enabled_for_loss", "source_workbook": "docs/method/value_basis/solar_wind_value_breakdown.xlsx",
                        "source_register_id": "WSW-S021", "crosswalk": rel(W_VAL), "unit": "2023_USD_per_kW",
                        "turbine_equipment_total": 1090, "other_direct_withheld": 239, "support_once": 294, "physical_reference": 1623,
                        "excluded_soft_sunk_nonphysical": 345, "installed_reference": 1968,
                        "reconciliation": "1090 + 239 + 294 = 1623; 1623 + 345 = 1968",
                        "site_override_required_for_loss": True, "values_embedded_as_curve_denominators": False},
        "hazard_axis": {"id": "WITHHELD_PENDING_QUALIFIED_ICE_RESPONSE_EVIDENCE", "status": "declared_not_active",
                        "runtime_damage_mapping_status": "withheld_no_runtime_curve",
                        "source_native_objects": [
                            {"id": "EQUIVALENT_RADIAL_GLAZE_ICE", "input_field": "ice_radial_thickness_mm", "unit": "mm",
                             "definition": "equivalent uniform radial glaze ice on a 25.4 mm cylinder at 10 m (ASCE 7 Ch.10 / Simple model)",
                             "role": "future_primary_axis", "runtime_curve_compatible": False, "source_id": "WSW-S020"},
                            {"id": "CONCURRENT_GUST", "input_field": "ice_concurrent_gust_3s_10m_mps", "unit": "m/s",
                             "role": "required_capture_for_deferred_ice_plus_wind", "runtime_curve_compatible": False, "source_id": "WSW-S011"},
                            {"id": "DESIGN_ICE", "input_field": "design_ice_thickness_mm", "unit": "mm",
                             "role": "future_overhead_line_denominator", "runtime_curve_compatible": False, "source_id": "WSW-S011"}]},
        "proposed_future_y_axis": {"id": "conditional_direct_repair_or_replacement_cost_ratio", "status": "withheld",
                                   "definition": "Expected direct repair-or-replacement cost divided by pre-event replacement value of the same failure unit given delivered ice state.",
                                   "denominator_rule": "Same failure unit in numerator and denominator; no whole-farm or installed-TIV substitution.",
                                   "support_cost_rule": "Allocate fieldwork and transport once after qualified disposition.",
                                   "excluded_from_ordinate": ["business_interruption", "icing_downtime", "derating_or_revenue", "frequency", "insurance_terms"]},
        "failure_units": [{"id": i, "subsystem": s, "component": c, "pathway": "ice_accretion", "treatment": t, "ordinate_status": "withheld",
                           "mechanism_note": m, "value_lineage": v} for i, s, c, t, m, v in WIND_UNITS],
        "candidate_evidence": {"status": "audit_only_not_runtime_shaped", "numeric_values_embedded": False, "runtime_enabled": False,
                               "candidate_families": ["overhead_line_field_withstand_1998", "distribution_pole_ice_wind_fragility_2024",
                                                      "lattice_tower_ice_fragility"],
                               "legacy_rejection": rel(wp("LEGACY_EVIDENCE_INGESTION_{cell}__{tag}.md"))},
        "curve_records": [],
        "parameter_tier_table": [],
        "derivation_rationale": {"decision": "NO_RUNTIME_CURVE",
                                 "reason": "No located evidence links ice thickness to physical repair cost for turbine or BOP units; the documented February 2021 icing produced outages without reported damage; the one damage case was falling ice; the overhead-line anchor lacks topology and a split value row.",
                                 "update_trigger": "Field, claims or work-order evidence at known ice thickness; Li et al. 2024 coefficients with site overhead share and split collection value."},
        "emit_contract": {"runtime_status": "withheld", "valid_input_behavior": "Return no numeric DR or loss and stable NO_RUNTIME_CURVE reason.",
                          "unknown_or_neighboring_pathway_behavior": "Reject or withhold; never default to ice_accretion.",
                          "annual_tail_boundary": "Frequency, EAL, PML, VaR and TVaR remain downstream and withheld."},
        "capability_declaration": wind_capability(),
    }


def wind_kats() -> dict:
    withheld = {m: {"value": None, "status": "withheld", "reason_codes": ["NO_RUNTIME_CURVE"]} for m in METRICS}
    tests = []
    cases = [
        ("WSW_VALID_ICE_19MM_BLADE_WITHHOLDS", "WT_BLADE_ASSEMBLY", {"ice_radial_thickness_mm": 19, "ice_concurrent_gust_3s_10m_mps": 15}),
        ("WSW_VALID_ICE_38MM_TOWER_WITHHOLDS", "WT_TOWER_AND_EXTERNAL_FIXTURES", {"ice_radial_thickness_mm": 38, "ice_concurrent_gust_3s_10m_mps": 25}),
        ("WSW_ZERO_ICE_WITHHOLDS_NOT_ZERO", "WT_NACELLE_AND_HUB", {"ice_radial_thickness_mm": 0}),
        ("WSW_OVERHEAD_LINE_WITH_DESIGN_STILL_WITHHOLDS", "WT_COLLECTION_NETWORK",
         {"ice_radial_thickness_mm": 45, "design_ice_thickness_mm": 13, "collection_topology": "overhead"}),
        ("WSW_MET_MAST_WITHHOLDS", "WT_CONTROL_AND_MET_STATION", {"ice_radial_thickness_mm": 30}),
        ("WSW_FOUNDATION_WITHHOLDS_NOT_ZERO", "WT_FOUNDATION", {"ice_radial_thickness_mm": 25}),
        ("WSW_LEGACY_TEXAS_CASE_NO_LEGACY_OUTPUT", "WT_NACELLE_AND_HUB", {"ice_radial_thickness_mm": 19.05}),
    ]
    for tid, unit, extra in cases:
        tests.append({"test_id": tid, "cell_id": "winter_storm_wind", "pathway_id": "ice_accretion", "failure_unit_id": unit,
                      "input": {"event_id": "e1", "event_family_id": "f1", "pathway_id": "ice_accretion", "failure_unit_id": unit, **extra},
                      "expected_output": withheld, "tolerance": None, "source": "fail_closed_contract",
                      "notes": "No numeric output is permitted in model v0.1."})
    rejects = [
        ("WSW_REJECT_MISSING_PATHWAY", {"event_id": "e1", "event_family_id": "f1", "failure_unit_id": "WT_BLADE_ASSEMBLY", "ice_radial_thickness_mm": 19}, "PATHWAY_ID_REQUIRED"),
        ("WSW_REJECT_DEFERRED_SHEDDING", {"event_id": "e1", "event_family_id": "f1", "pathway_id": "ice_shedding_impact", "failure_unit_id": "WT_BLADE_ASSEMBLY"}, "PATHWAY_DEFERRED_NOT_IMPLEMENTED"),
        ("WSW_REJECT_MISSING_EVENT_FAMILY", {"event_id": "e1", "pathway_id": "ice_accretion", "failure_unit_id": "WT_BLADE_ASSEMBLY", "ice_radial_thickness_mm": 19}, "EVENT_FAMILY_ID_REQUIRED"),
        ("WSW_REJECT_UNKNOWN_UNIT", {"event_id": "e1", "event_family_id": "f1", "pathway_id": "ice_accretion", "failure_unit_id": "ROTOR"}, "FAILURE_UNIT_ID_UNKNOWN"),
    ]
    for tid, payload, code in rejects:
        tests.append({"test_id": tid, "cell_id": "winter_storm_wind", "pathway_id": payload.get("pathway_id"), "failure_unit_id": payload.get("failure_unit_id"),
                      "input": payload, "expected_error_code": code, "tolerance": None, "source": "fail_closed_contract",
                      "notes": "Contract rejection precedes withholding."})
    return {"schema_version": "known_answer_tests.v1", "cell_id": "winter_storm_wind", "semantic_damage_model_version": "model v0.1",
            "documentation_revision": "docs r1", "runtime_curve_count": 0, "runtime_curve_known_answer_tests": 0, "tests": tests}


# ================================================================================ shared writers
def claims_rows(claims) -> list[dict]:
    rows = []
    for cid, ctype, srcs, loc, tier, param, status, permitted, prohibited, reasoning, trigger, text in claims:
        pathway = "snow_gravity_load" if cid.startswith("WSS") else "ice_accretion"
        rows.append(dict(zip(R.CLAIM_FIELDS, [cid, pathway, text, ctype, srcs, loc, tier, param, status, permitted, prohibited, reasoning, trigger])))
    return rows


def value_rows(rows, source_id, sheet, unit, pathway) -> list[dict]:
    out = []
    for row, rid, label, value, fclass, unit_id, role, include, rule, guard, status in rows:
        out.append(dict(zip(R.VALUE_FIELDS, [source_id, f"docs/method/value_basis/solar_wind_value_breakdown.xlsx::{sheet}!{row}", rid, label,
                                             repr(float(value)) if isinstance(value, float) else str(value), unit, fclass, unit_id, pathway,
                                             role, include, rule, guard, status, ""])))
    return out


def workbook(path: Path, title: str, sheets: list[tuple[str, list[str], list[list]]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
    wb.properties.title = title
    wb.properties.creator = "owner-directed (JL) - agent-drafted"
    wb.save(path)


def build_solar() -> None:
    SOLAR_DIR.mkdir(parents=True, exist_ok=True)
    art = solar_artifact()
    write_json(S_ART, art)
    write_json(S_CAP, art["capability_declaration"])
    kats = solar_kats()
    write_json(S_KAT, kats)
    write_csv(S_SRC, R.SOURCE_FIELDS, R.SOLAR_SOURCES)
    claims = claims_rows(R.SOLAR_CLAIMS)
    write_csv(S_CLM, R.CLAIM_FIELDS, claims)
    params = [dict(zip(R.PARAMETER_FIELDS, p)) for p in R.SOLAR_PARAMETERS]
    write_csv(S_PAR, R.PARAMETER_FIELDS, params)
    values = value_rows(R.SOLAR_VALUE_ROWS, "NLR_Q1_2025_UPV_PV_ONLY_2024_USD", "Solar_Map", "2024_USD_per_kWdc", "snow_gravity_load")
    write_csv(S_VAL, R.VALUE_FIELDS, values)

    r_grid = [round(0.1 * i, 2) for i in range(0, 41)]
    q_grid = [round(0.2 * i, 2) for i in range(0, 31)]
    curves = [[r] + [struct_drs(r)[s] for s in STRUCT_SCEN] for r in r_grid]
    mod = [[q] + [_dr(q, MOD_SCEN[s], MOD_COST, MODULE_BETA) for s in MOD_SCEN] for q in q_grid]
    kat_rows = [[t["test_id"], t["failure_unit_id"], t["expected"]["axis_ratio"], t["expected"].get("axis_array_load_kpa"),
                 t["expected"]["scenario_drs"]["lower_resistance"], t["expected"]["scalar_central_dr"], t["expected"]["scenario_drs"]["upper_resistance"]]
                for t in kats["runtime_curve_known_answer_tests"]]
    workbook(S_WB, "winter_storm_solar model v0.1 docs r1 derivation and audit view", [
        ("README", ["field", "value"], [["cell_id", "winter_storm_solar"], ["model", "model v0.1 / docs r1 (proposed, not indexed)"],
                                        ["grade", "screening_engineering_proxy"], ["runtime artifact", rel(S_ART)],
                                        ["note", "Derivation/audit view only; the JSON artifact is the runtime object."]]),
        ("Structure_Curve", ["event_to_design_ratio", "DR_lower_resistance", "DR_central_screening", "DR_upper_resistance"], curves),
        ("Module_Direct_Curve", ["array_snow_load_kpa", "direct_DR_lower", "direct_DR_central", "direct_DR_upper"], mod),
        ("Rack_Tests", ["source_id", "specimen", "design_pa", "failure_pa", "failure_over_design"],
         [["WSS-S022", "i steel 4x2", 1294, 4250, 3.28], ["WSS-S022", "ii steel 4x4", 1294, 1750, 1.35],
          ["WSS-S022", "iii aluminium 4x5", 1036, 1000, 0.97], ["WSS-S022", "iv aluminium+steel 4x5", 1036, 2250, 2.17],
          ["WSS-S022", "v aluminium+steel 4x6", 688, 2000, 2.91], ["WSS-S022", "vi steel scaffold pipe 4x4", 1036, 1250, 1.21],
          ["WSS-S023", "JSME thin-wall steel truss", 1980, 5038, round(5038 / 1980, 4)]]),
        ("Value_Crosswalk", R.VALUE_FIELDS, [[r[f] for f in R.VALUE_FIELDS] for r in values]),
        ("Parameter_Tiers", R.PARAMETER_FIELDS, [[r[f] for f in R.PARAMETER_FIELDS] for r in params]),
        ("Claim_Register", R.CLAIM_FIELDS, [[r[f] for f in R.CLAIM_FIELDS] for r in claims]),
        ("Source_Register", R.SOURCE_FIELDS, [[r[f] for f in R.SOURCE_FIELDS] for r in R.SOLAR_SOURCES]),
        ("Known_Answers", ["test_id", "failure_unit_id", "ratio", "array_load_kpa", "DR_lower", "DR_central", "DR_upper"], kat_rows),
        ("QA_Checks", ["check", "status"], [["artifact schema v3", "see validator"], ["value reconciliation 877.80 / 1120", "see validator"],
                                            ["known answers independent of evaluator", "see validator"]]),
    ])


def build_wind() -> None:
    WIND_DIR.mkdir(parents=True, exist_ok=True)
    art = wind_artifact()
    write_json(W_ART, art)
    write_json(W_CAP, art["capability_declaration"])
    write_json(W_KAT, wind_kats())
    write_csv(W_SRC, R.SOURCE_FIELDS, R.WIND_SOURCES)
    claims = claims_rows(R.WIND_CLAIMS)
    write_csv(W_CLM, R.CLAIM_FIELDS, claims)
    params = [dict(zip(R.PARAMETER_FIELDS, p)) for p in R.WIND_PARAMETERS]
    write_csv(W_PAR, R.PARAMETER_FIELDS, params)
    values = value_rows(R.WIND_VALUE_ROWS, "NREL_CWER_2024", "Wind_Map", "2023_USD_per_kW", "ice_accretion")
    write_csv(W_VAL, R.VALUE_FIELDS, values)
    legacy = [["ROTOR_ASSEMBLY", 0.60, 3.5, 1.50], ["TOWER", 0.30, 2.0, 2.50], ["NACELLE", 0.35, 3.0, 0.75], ["ELECTRICAL", 0.70, 4.0, 0.75]]
    audit = []
    for name, L, k, x0 in legacy:
        f = lambda x: L / (1 + math.exp(-k * (x - x0)))
        audit.append([name, L, k, x0, f(0.0), f(0.75), f(1.5), "rejected"])
    workbook(W_WB, "winter_storm_wind model v0.1 docs r1 scaffold audit view", [
        ("README", ["field", "value"], [["cell_id", "winter_storm_wind"], ["model", "model v0.1 / docs r1 fail-closed scaffold"],
                                        ["runtime", "NO_RUNTIME_CURVE: every DR and loss withheld"]]),
        ("Failure_Units", ["id", "subsystem", "component", "treatment", "mechanism_note", "value_lineage"], [list(u) for u in WIND_UNITS]),
        ("Legacy_Audit", ["legacy_unit", "L", "k", "x0_in", "DR_at_0_in", "DR_at_0.75_in", "DR_at_1.5_in", "decision"], audit),
        ("Value_Crosswalk", R.VALUE_FIELDS, [[r[f] for f in R.VALUE_FIELDS] for r in values]),
        ("Parameter_Tiers", R.PARAMETER_FIELDS, [[r[f] for f in R.PARAMETER_FIELDS] for r in params]),
        ("Claim_Register", R.CLAIM_FIELDS, [[r[f] for f in R.CLAIM_FIELDS] for r in claims]),
        ("Source_Register", R.SOURCE_FIELDS, [[r[f] for f in R.SOURCE_FIELDS] for r in R.WIND_SOURCES]),
        ("QA_Checks", ["check", "status"], [["curve_records empty", "see validator"], ["all metrics withheld", "see validator"]]),
    ])


if __name__ == "__main__":
    build_solar()
    build_wind()
    print("built", rel(SOLAR_DIR), "and", rel(WIND_DIR))
