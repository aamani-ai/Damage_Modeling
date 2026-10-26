"""Validator for the noncanonical tropical-cyclone wind × real estate model-v0.1 / docs-r1 package.

Checks: required files; JSON schemas (bundle v3, capability v3, damage emit v2, when jsonschema is
installed); embedded vs standalone capability; rectangular registers and resolvable source IDs; artifact
ordinates vs the governed ordinates CSV (values, Hazus wbIDs, ranges, dips); loader compatibility (selector
routing unique, one pathway range); builder reproducibility of the JSON outputs; every known-answer test
through the reference evaluator; non-promotion (absent from the artifact index and version registry, no
current/ package); workbook sheets; links in the cell's Markdown.

Usage (from the repository root):
    python scripts/reference_helpers/validate_tropical_cyclone_wind_real_estate_v0_1_package.py
Exit code 0 when every check passes.
"""
from __future__ import annotations

import csv
import json
import math
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from tropical_cyclone_wind_real_estate_curve_eval import DamageCallError, evaluate_damage_call  # noqa: E402
import build_tropical_cyclone_wind_real_estate_v0_1_package as builder  # noqa: E402

CELL_ID = builder.CELL_ID
CELL = ROOT / f"docs/cells/{CELL_ID}"
PROP = CELL / "proposed"
TAG = builder.TAG
SCHEMAS = ROOT / "docs/contracts/schemas"
results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((name, bool(ok), detail))


def load(path: Path):
    with open(path) as fh:
        return json.load(fh)


def csv_rows(path: Path) -> list[list[str]]:
    with open(path, newline="") as fh:
        return list(csv.reader(fh))


def main() -> int:
    files = {
        "artifact": PROP / f"{TAG}__curve_artifact.json",
        "capability": PROP / f"{TAG}__capability.json",
        "kats": PROP / f"known_answer_tests_{TAG}.json",
        "workbook": PROP / f"damage_curve_records_{TAG}.xlsx",
        "dossier": PROP / f"{CELL_ID}_curve_derivation_dossier__model_v0_1__docs_r1.md",
        "metadata_spec": PROP / f"{CELL_ID}_damage_code_metadata_spec__model_v0_1__docs_r1.md",
        "validation_report": PROP / f"VALIDATION_REPORT_{TAG}.md",
        "source_register": PROP / f"SOURCE_REGISTER_{TAG}.csv",
        "claim_register": PROP / f"CLAIM_PARAMETER_REGISTER_{TAG}.csv",
        "tiers": PROP / f"PARAMETER_TIER_TABLE_{TAG}.csv",
        "crosswalk": PROP / f"VALUE_CROSSWALK_{TAG}.csv",
        "ordinates": PROP / f"CURVE_ORDINATES_{TAG}.csv",
        "sensitivity": PROP / f"VARIANT_SENSITIVITY_{TAG}.csv",
    }
    missing = [k for k, p in files.items() if not p.exists()]
    check("required files exist", not missing, ", ".join(missing))
    if missing:
        return report()

    art, cap, kats = load(files["artifact"]), load(files["capability"]), load(files["kats"])
    pathway = art["pathways"][0]
    recs = pathway["curve_records"]
    check("embedded capability equals standalone capability", art["capability_declaration"] == cap)
    for key in ("source_dossier", "source_workbook", "known_answer_tests", "source_register",
                "claim_parameter_register", "value_crosswalk"):
        check(f"artifact path resolves: {key}", (ROOT / art[key]).exists(), art[key])

    # schemas
    try:
        from jsonschema import Draft202012Validator
        from referencing import Registry, Resource
        sch = {n: load(SCHEMAS / f"{n}.schema.json") for n in
               ("curve_artifact_bundle.v3", "capability_declaration.v3", "damage_emit.v2")}
        reg = Registry().with_resources([(s["$id"], Resource.from_contents(s)) for s in sch.values()])
        validators = {n: Draft202012Validator(s, registry=reg) for n, s in sch.items()}
        errs = list(validators["curve_artifact_bundle.v3"].iter_errors(art))
        check("artifact validates against bundle v3", not errs, "; ".join(e.message[:120] for e in errs[:3]))
        errs = list(validators["capability_declaration.v3"].iter_errors(cap))
        check("capability validates against capability v3", not errs, "; ".join(e.message[:120] for e in errs[:3]))
        emit_validator = validators["damage_emit.v2"]
    except ModuleNotFoundError:
        check("jsonschema available for schema checks", False, "install jsonschema>=4.17")
        emit_validator = None

    # registers
    for k in ("source_register", "claim_register", "tiers", "crosswalk", "ordinates", "sensitivity"):
        widths = {len(r) for r in csv_rows(files[k])}
        check(f"register rectangular: {files[k].name}", len(widths) == 1, str(widths))
        check(f"register uses LF line endings: {files[k].name}", b"\r" not in files[k].read_bytes())
    sources = {r[0] for r in csv_rows(files["source_register"])[1:]}
    refs = {ref.split("#")[0] for r in recs for ref in r["source_parameter_refs"]}
    check("curve source refs resolve to the source register", refs <= sources, str(refs - sources))
    for k, col in (("claim_register", "source_ids"), ("tiers", "source_ids")):
        ids = set()
        for r in csv.DictReader(open(files[k])):
            ids |= {s for s in r[col].split(";") if s.startswith("TCRE-S")}
        check(f"{files[k].name.split('_tropical')[0]} source IDs resolve", ids <= sources, str(ids - sources))
    ord_src = {s for r in csv.DictReader(open(files["ordinates"])) for s in r["source_ids"].split(";")}
    check("ordinate source IDs resolve", ord_src <= sources, str(ord_src - sources))

    # ordinates and records
    pts = builder.ordinates()
    ocsv = list(csv.DictReader(open(files["ordinates"])))
    mism = []
    for r in recs:
        key = (r["selector_match"]["reference_building_id"], r["selector_match"]["hazus_terrain_class"])
        if r["parameters"]["points"] != pts[key]:
            mism.append(r["curve_id"])
        spec = builder.BUILDINGS[key[0]]
        wbids = {int(o["hazus_wbid"]) for o in ocsv if o["reference_building_id"] == key[0]}
        sbts = {o["hazus_sbt"] for o in ocsv if o["reference_building_id"] == key[0]}
        if wbids != {spec["wbid"]} or sbts != {spec["sbt"]} or pathway["hazus_variants"][key[0]]["wbid"] != spec["wbid"] \
                or f"wbID={spec['wbid']};" not in r["source_parameter_refs"][1]:
            mism.append(r["curve_id"] + ":wbid")
    check("artifact ordinates and Hazus variants equal the governed CSV (15 records)", not mism and len(recs) == 15, str(mism))
    speeds = [float(s) for s in range(50, 255, 5)]
    check("every record has the 41 published speeds 50..250 mph",
          all([p[0] for p in r["parameters"]["points"]] == speeds for r in recs))
    check("every value is within [0, 1]", all(0.0 <= p[1] <= 1.0 for r in recs for p in r["parameters"]["points"]))
    check("every record is 0 at 50 mph (supports the below-range zero)", all(r["parameters"]["points"][0][1] == 0.0 for r in recs))
    bad_dips = [(r["curve_id"], a[0]) for r in recs for a, b in zip(r["parameters"]["points"], r["parameters"]["points"][1:])
                if b[1] < a[1] and (a[1] - b[1] > builder.DIP_LIMIT or b[0] <= 200)]
    check(f"dips are at most {builder.DIP_LIMIT} and only above 200 mph", not bad_dips, str(bad_dips[:5]))
    flagged = {r["curve_id"] for r in recs if "HAZUS_SIMULATION_NOISE_SMALL_DIPS" in r["metadata_flags"]}
    dipping = {r["curve_id"] for r in recs if builder.has_dips(r["parameters"]["points"])}
    check("records with dips carry the dip flag (and only they)", flagged == dipping, str(flagged ^ dipping))

    # loader compatibility (Hazard generic loader: unique selector match, one pathway range)
    sel_fields = pathway["record_routing_fields"]
    required_sel = [s["field"] for s in pathway["selector_logic"] if s["required"] is True]
    check("routing fields are exactly the required selectors", sel_fields == required_sel, str(required_sel))
    combos = [tuple(r["selector_match"][f] for f in sel_fields) for r in recs]
    check("selector_match keys equal the routing fields", all(set(r["selector_match"]) == set(sel_fields) for r in recs))
    char_sel = {s["field"]: s["allowed"] for s in pathway["selector_logic"] if s["required"] is False}
    refs = {}
    for b in pathway["wind_characteristic_contract"].values():
        for f, c in b.items():
            refs.setdefault(f, set()).add(c["reference"])
    check("optional characteristic selectors admit only reference values",
          set(char_sel) == set(refs) and all(set(v) == refs[f] for f, v in char_sel.items()))
    check("every building × terrain routes to exactly one record", len(set(combos)) == len(combos) == 15)
    check("all records share the pathway valid_range", all(r["valid_range"] == pathway["hazard_axis"]["valid_range"] for r in recs))
    check("no conditioner is required", pathway["conditioner_logic"] == [] and art["evaluation_contract"]["required_conditioner_fields"] == [])

    # reproducibility of JSON outputs
    check("builder reproduces the artifact JSON", builder.artifact(pts) == art)
    check("builder reproduces the known-answer tests", builder.known_answer_tests(pts) == kats)

    # known-answer tests
    tol = kats["absolute_tolerance"]
    fails = []
    n = 0
    for t in kats["formula_known_answer_tests"] + kats["withheld_tests"]:
        n += 1
        try:
            out = evaluate_damage_call(art, t["input"])
        except DamageCallError as e:
            fails.append(f"{t['test_id']}: error {e.code}")
            continue
        if emit_validator is not None:
            emit_errs = list(emit_validator.iter_errors(out))
            if emit_errs:
                fails.append(f"{t['test_id']}: emit schema {emit_errs[0].message[:80]}")
        res, exp = out["failure_unit_results"][0], t["expected"]
        if res["status"] != exp["status"]:
            fails.append(f"{t['test_id']}: status {res['status']}")
            continue
        if exp["status"] == "withheld":
            if exp["reason_code"] not in res["withheld_reason_codes"] or res["scalar_central_dr"] is not None:
                fails.append(f"{t['test_id']}: reasons {res['withheld_reason_codes']}")
            continue
        if res["curve_id"] != exp["curve_id"] or not math.isclose(res["scalar_central_dr"], exp["failure_unit_damage_ratio"], rel_tol=0, abs_tol=tol):
            fails.append(f"{t['test_id']}: {res['curve_id']} {res['scalar_central_dr']}")
        for f in exp.get("metadata_flags_include", []):
            if f not in res["metadata_flags"]:
                fails.append(f"{t['test_id']}: missing flag {f}")
    for t in kats["error_tests"]:
        n += 1
        try:
            evaluate_damage_call(art, t["input"])
            fails.append(f"{t['test_id']}: no error raised")
        except DamageCallError as e:
            if e.code != t["expected_error_code"]:
                fails.append(f"{t['test_id']}: {e.code}")
    check(f"known-answer tests pass ({n})", not fails, "; ".join(fails[:5]))

    # non-promotion
    index = (ROOT / "docs/contracts/machine_readable_artifact_index.json").read_text()
    registry = (ROOT / "docs/cells/VERSION_REGISTRY.md").read_text()
    check("proposal absent from the machine-readable artifact index", f"{TAG}__curve_artifact.json" not in index)
    check("proposal pin absent from the version registry", f"{CELL_ID}@model_v0_1__docs_r1" not in registry)
    check("proposal is not installed as current", not (CELL / "current" / f"{TAG}__curve_artifact.json").exists())
    check("artifact is noncanonical and proposed",
          art["canonical_runtime_artifact"] is False and art["schema_status"] == "proposed_draft"
          and art["promotion_status"] == "proposed" and cap["promotion_gate"]["status"] == "blocked")

    # workbook
    try:
        from openpyxl import load_workbook
        wb = load_workbook(files["workbook"], read_only=True)
        need = {"README", "curve_records", "reference_buildings", "wind_characteristics", "sources", "claims",
                "parameter_tiers", "value_crosswalk", "variant_sensitivity", "known_answer_tests"}
        check("workbook has the expected sheets", need <= set(wb.sheetnames), str(need - set(wb.sheetnames)))
    except ModuleNotFoundError:
        check("openpyxl available for workbook check", False, "install openpyxl")

    # links
    broken = []
    for md in CELL.rglob("*.md"):
        for m in re.finditer(r"\]\(([^)#]+)(#[^)]*)?\)", md.read_text()):
            target = m.group(1)
            if not target.startswith("http") and not (md.parent / target).exists():
                broken.append(f"{md.name} -> {target}")
    check("Markdown links in the cell resolve", not broken, "; ".join(broken[:5]))
    return report()


def report() -> int:
    width = max(len(n) for n, _, _ in results)
    for name, ok, detail in results:
        print(f"{'PASS' if ok else 'FAIL'}  {name.ljust(width)}  {detail if not ok else ''}")
    failed = sum(1 for _, ok, _ in results if not ok)
    print(f"\n{len(results) - failed}/{len(results)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
